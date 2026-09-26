"""
inventory.py
------------
Material/thickness-aware stock management for multi-part nesting jobs.

A real shop doesn't have "one infinite sheet size" -- it has a stock list
(standard sheets it can buy/has on hand) plus remnants (offcuts from past
jobs), each tagged with material + thickness, and a part can only be nested
onto stock of its own exact material + thickness (you cannot cut a 3mm
mild-steel bracket and a 5mm stainless bracket from the same sheet). This
module:
  1. Groups an assembly's parts by (material, thickness) -- each group is
     its own independent nesting job, never mixed with another.
  2. Matches each group to its best-fit stock sheet from an inventory list
     (remnants preferred over full sheets, to burn down scrap first; then
     the smallest matching size that's still big enough for the group's
     largest part).
  3. Runs one `Nester` job per group on that stock size, then reconciles
     against how many sheets of that stock are actually on hand --
     "geometrically these parts fit on 6 sheets" is different from "you
     only have 4 sheets of this stock in inventory," and the two get
     reported separately.

Within a group, stock is consumed size-by-size, best match first (remnants
before full sheets): nest onto the current best-matching stock, and
whatever doesn't fit -- either geometrically, or because that stock's
on-hand quantity ran out -- cascades to the next best-matching stock still
in the inventory, and so on until either everything is placed or no
matching stock is left.

By default, this cascades one stock size to the next *between* nesting
passes, but each individual pass still nests onto a single sheet size --
it's not a joint optimization across sizes within one pass (e.g. it can't
jointly weigh "would 2 sheets of A + 1 of B beat 3 sheets of B" the way a
real joint solver could). `run_job()`'s `true_joint_stock_optimization`
flag closes that gap via `stock_solver.py`'s heuristic search over an
explicit, ordered plan of specific stock picks (possibly mixed sizes) --
see that module's docstring for the search and its honest cost.
"""

import datetime
import json
import os
import uuid
from dataclasses import dataclass, field
from typing import Dict, List, Optional, Tuple

from genetic import optimize_order
from geometry import polygon_bbox
from nester import Nester, Part, PlacedPart
from remnant import largest_empty_rect


@dataclass
class StockSheet:
    material: str
    thickness: float
    width: float
    height: float
    quantity: Optional[int] = None  # None = unlimited (e.g. a standard sheet you can always reorder)
    is_remnant: bool = False
    id: Optional[str] = None
    # Sheet metal is bought and priced by WEIGHT, not by the sheet -- so
    # cost is derived from density + dimensions rather than entered
    # directly. Both None = unpriced (job-cost estimates degrade to "n/a"
    # for this entry, same as the old price_per_sheet's None did).
    price_per_kg: Optional[float] = None    # material cost, currency/kg -- whatever you enter,
                                             # not a weight/alloy-derived guess
    density_g_cm3: Optional[float] = None   # e.g. 7.85 mild steel, 8.0 stainless, 2.70 aluminum --
                                             # needed to turn this entry's own width/height/thickness
                                             # into a weight; None = this entry's weight/cost can't be
                                             # computed at all, regardless of price_per_kg
    scrap_price_per_kg: Optional[float] = None  # what THIS material's offcuts sell for/kg -- per
                                                 # entry, not a single blended rate, since scrap value
                                                 # genuinely varies a lot by material (aluminum scrap is
                                                 # worth much more per kg than mild steel, for example).
                                                 # None = this entry's scrap isn't valued in reports.

    def area(self):
        return self.width * self.height

    def weight_kg(self) -> Optional[float]:
        """This exact sheet's full weight (width x height x thickness x
        density_g_cm3, mm and g/cm^3 in, kg out) -- None if density_g_cm3
        isn't set."""
        if self.density_g_cm3 is None:
            return None
        return self.width * self.height * self.thickness * self.density_g_cm3 / 1e6

    def material_cost(self) -> Optional[float]:
        """What buying/consuming ONE of this exact stock sheet costs (its
        full weight x price_per_kg) -- None if either weight or
        price_per_kg is unknown. For a remnant, this is what it WOULD have
        cost to buy that same amount of material new -- i.e. "money saved"
        by using it, not cash actually spent this job (see nesting_widgets.py's
        job-financials calculation)."""
        weight = self.weight_kg()
        if weight is None or self.price_per_kg is None:
            return None
        return weight * self.price_per_kg


# ------------------------------------------------------------------ file I/O

# The stock list is a TABLE -- one row per stock entry, the same columns the
# Stock tab shows -- so its file format is a spreadsheet, not JSON. A .xlsx
# is something the person who actually owns the stock can open, edit, sort,
# filter and send on without going near a text editor or getting a comma
# wrong; a .json is not. Excel is therefore the format `save_inventory()`
# writes and the one new inventories should use.
#
# JSON is still READ (and still written back when the loaded path is a
# .json), so inventories written before this change -- and anything a script
# generates -- keep working untouched. `convert_inventory()` moves one over.
EXCEL_EXTENSIONS = (".xlsx", ".xlsm")
DEFAULT_INVENTORY_EXTENSION = ".xlsx"
STOCK_SHEET_NAME = "Stock"


def _is_excel_path(path) -> bool:
    return os.path.splitext(str(path))[1].lower() in EXCEL_EXTENSIONS


def _require_openpyxl():
    """openpyxl is imported lazily, and only by the .xlsx paths: this module
    is also imported inside FreeCAD's bundled Python, where an install can't
    be assumed, and nothing but Excel I/O needs it there."""
    try:
        import openpyxl
    except ImportError:
        raise RuntimeError(
            "Reading/writing an Excel stock file needs the openpyxl package "
            "(pip install openpyxl)."
        ) from None
    return openpyxl


# --- cell parsing -----------------------------------------------------------
#
# Every parser below takes a raw cell value, which -- because a human types
# into these -- may be a real number, a string that looks like one, or blank.
# Blank means "not set" everywhere: unlimited quantity, unpriced stock.

_TRUE_WORDS = {"1", "y", "yes", "true", "t", "x", "remnant", "offcut"}
_FALSE_WORDS = {"0", "n", "no", "false", "f", "sheet", "full", "full sheet", "new"}


def _clean(value) -> str:
    """A cell as trimmed text. Thousands separators are dropped so a width
    typed as `1,220` reads the same as `1220`."""
    if value is None:
        return ""
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    return str(value).strip().replace(",", "")


def _parse_optional_text(value):
    text = _clean(value)
    return text or None


def _parse_text(value) -> str:
    return _clean(value)


def _parse_float(value) -> float:
    text = _clean(value)
    return float(text) if text else 0.0


def _parse_optional_float(value):
    text = _clean(value)
    return float(text) if text else None


def _parse_optional_int(value):
    text = _clean(value)
    if not text:
        return None
    return int(float(text))  # `3.0` from a spreadsheet cell is still 3 sheets


def _parse_bool(value) -> bool:
    if isinstance(value, bool):
        return value
    text = _clean(value).lower()
    if text in _TRUE_WORDS:
        return True
    if text in _FALSE_WORDS or not text:
        return False
    raise ValueError(f"expected Yes or No, got {text!r}")


def _format_bool(value):
    return "Yes" if value else "No"


def _format_plain(value):
    return value


# field name -> (column heading, parse cell -> value, format value -> cell).
# This list IS the sheet layout: column order, headings and the Stock tab's
# own columns all come from here, so the spreadsheet and the in-app table
# can never drift apart.
_STOCK_FIELDS = [
    ("id", "ID", _parse_optional_text, _format_plain),
    ("material", "Material", _parse_text, _format_plain),
    ("thickness", "Thickness (mm)", _parse_float, _format_plain),
    ("width", "Width (mm)", _parse_float, _format_plain),
    ("height", "Height (mm)", _parse_float, _format_plain),
    ("quantity", "Quantity (blank=unlimited)", _parse_optional_int, _format_plain),
    ("is_remnant", "Remnant", _parse_bool, _format_bool),
    ("price_per_kg", "Price/kg (blank=unpriced)", _parse_optional_float, _format_plain),
    ("density_g_cm3", "Density g/cm³ (blank=unpriced)", _parse_optional_float, _format_plain),
    ("scrap_price_per_kg", "Scrap price/kg (blank=unpriced)", _parse_optional_float, _format_plain),
]

STOCK_COLUMNS = [heading for _f, heading, _p, _fmt in _STOCK_FIELDS]
# The same columns, shortened for the app's own table. The parenthesised
# hint exists to make the SPREADSHEET self-explanatory -- a workbook has
# nowhere else to say "blank means unlimited" -- but in the UI that's what
# a tooltip is for, and a heading long enough to truncate says nothing.
STOCK_COLUMN_LABELS = [heading.split("(")[0].strip() for heading in STOCK_COLUMNS]
_PARSERS = {field: parse for field, _h, parse, _fmt in _STOCK_FIELDS}

# Column headings people plausibly type instead of the canonical ones. Only
# genuinely different WORDS need to live here -- spelling, case, spacing and
# the parenthesised units/hints are all handled by _normalize_heading().
_HEADING_ALIASES = {
    "stockid": "id",
    "sheetid": "id",
    "qty": "quantity",
    "count": "quantity",
    "onhand": "quantity",
    "isremnant": "is_remnant",
    "offcut": "is_remnant",
    "materialprice": "price_per_kg",
}


def _normalize_heading(text) -> str:
    """`Thickness (mm)`, `thickness` and `THICKNESS` are the same column:
    units and hints live in parentheses, so everything from the first `(` on
    is dropped, then anything that isn't a letter or digit."""
    head = str(text or "").split("(")[0].strip().lower()
    return "".join(ch for ch in head if ch.isalnum())


def _column_map(heading_row):
    """{column index: field name} for one row of headings. Unrecognized
    columns are ignored rather than fatal -- a shop that adds its own
    `Supplier` or `Bin` column to the sheet keeps it, and we keep reading
    the columns we know."""
    canonical = {_normalize_heading(h): f for f, h, _p, _fmt in _STOCK_FIELDS}
    mapping = {}
    for index, cell in enumerate(heading_row):
        norm = _normalize_heading(cell)
        if not norm:
            continue
        field = canonical.get(norm) or _HEADING_ALIASES.get(norm)
        if field is None:
            # `Density` for `Density g/cm³`, `Price` for `Price/kg` -- a
            # heading shortened (or lengthened) at the end still matches.
            for canon_norm, canon_field in canonical.items():
                if canon_norm.startswith(norm) or norm.startswith(canon_norm):
                    field = canon_field
                    break
        if field is not None and field not in mapping.values():
            mapping[index] = field
    return mapping


def _find_heading_row(rows):
    """Index of the row that holds the column headings, or None. Searched
    for (over the first 20 rows) rather than assumed to be row 1, so a sheet
    carrying a title or a note above the table still loads."""
    best, best_score = None, 0
    for index, row in enumerate(rows[:20]):
        score = len(_column_map(row))
        if score > best_score:
            best, best_score = index, score
    return best if best_score >= 3 else None


def _rows_to_stock(rows, heading_index, source) -> List[StockSheet]:
    columns = _column_map(rows[heading_index])
    stock = []
    for offset, row in enumerate(rows[heading_index + 1:], start=heading_index + 2):
        values = {"material": "", "thickness": 0.0, "width": 0.0, "height": 0.0}
        for index, field in columns.items():
            raw = row[index] if index < len(row) else None
            try:
                values[field] = _PARSERS[field](raw)
            except ValueError as e:
                raise ValueError(
                    f"{source} row {offset}, column "
                    f"{rows[heading_index][index]!r}: {e}"
                ) from None
        # A row is only an entry if it says WHAT it is -- blank spacer rows,
        # and a trailing row holding nothing but a stray note, are skipped.
        if not values.get("material") and not values.get("id"):
            continue
        stock.append(StockSheet(**values))
    return stock


def load_inventory_xlsx(path) -> List[StockSheet]:
    """Read the stock table from an Excel workbook: the `Stock` sheet if
    there is one, otherwise the first sheet."""
    openpyxl = _require_openpyxl()
    # data_only: a cell computed by a formula reads as its last value saved
    # by Excel (a formula never saved by Excel has no cached value and reads
    # as blank -- open and save the file in Excel once and it fills in).
    workbook = openpyxl.load_workbook(path, data_only=True, read_only=True)
    try:
        if STOCK_SHEET_NAME in workbook.sheetnames:
            worksheet = workbook[STOCK_SHEET_NAME]
        else:
            worksheet = workbook.worksheets[0]
        rows = [list(r) for r in worksheet.iter_rows(values_only=True)]
    finally:
        workbook.close()

    heading_index = _find_heading_row(rows)
    if heading_index is None:
        raise ValueError(
            f"{os.path.basename(str(path))} has no stock table: expected a heading row "
            f"with columns like {', '.join(STOCK_COLUMNS[:5])}."
        )
    return _rows_to_stock(rows, heading_index, os.path.basename(str(path)))


def save_inventory_xlsx(path, stock: List[StockSheet]):
    """Write the stock table to an Excel workbook, styled to be worked in:
    frozen, bold headings, a filter row, Yes/No validation on Remnant.

    Any OTHER sheet in an existing workbook is left alone -- a shop's own
    notes/pricing tab lives on across the rewrite that every committed job
    triggers (see `commit_job()`); only the `Stock` sheet is ours."""
    openpyxl = _require_openpyxl()
    from openpyxl.styles import Alignment, Font
    from openpyxl.utils import get_column_letter
    from openpyxl.worksheet.datavalidation import DataValidation

    workbook, index = None, 0
    if os.path.exists(path):
        try:
            workbook = openpyxl.load_workbook(path)
        except Exception:
            workbook = None  # not a workbook we can read -- replace it wholesale
    if workbook is None:
        workbook = openpyxl.Workbook()
        workbook.remove(workbook.active)
    elif STOCK_SHEET_NAME in workbook.sheetnames:
        previous = workbook[STOCK_SHEET_NAME]
        index = workbook.index(previous)
        workbook.remove(previous)
    worksheet = workbook.create_sheet(STOCK_SHEET_NAME, index)

    worksheet.append(STOCK_COLUMNS)
    for s in stock:
        worksheet.append([fmt(getattr(s, field)) for field, _h, _p, fmt in _STOCK_FIELDS])

    for column, heading in enumerate(STOCK_COLUMNS, start=1):
        cell = worksheet.cell(row=1, column=column)
        cell.font = Font(bold=True)
        cell.alignment = Alignment(vertical="center", wrap_text=True)
        worksheet.column_dimensions[get_column_letter(column)].width = min(
            26, max(11, len(heading) + 2)
        )
    worksheet.freeze_panes = "A2"
    last_row = max(2, worksheet.max_row)
    worksheet.auto_filter.ref = f"A1:{get_column_letter(len(STOCK_COLUMNS))}{last_row}"

    remnant_column = get_column_letter(
        [f for f, _h, _p, _fmt in _STOCK_FIELDS].index("is_remnant") + 1
    )
    remnant_choices = DataValidation(type="list", formula1='"Yes,No"', allow_blank=True)
    worksheet.add_data_validation(remnant_choices)
    # A generous range, not just the rows written: the dropdown should be
    # there for the rows the user is about to ADD, which is the whole point.
    remnant_choices.add(f"{remnant_column}2:{remnant_column}{last_row + 200}")

    workbook.save(path)


def load_inventory_json(path) -> List[StockSheet]:
    with open(path, encoding="utf-8") as f:
        data = json.load(f)
    return [StockSheet(**s) for s in data["stock"]]


def save_inventory_json(path, stock: List[StockSheet]):
    payload = {"stock": [vars(s) for s in stock]}
    with open(path, "w", encoding="utf-8") as f:
        json.dump(payload, f, indent=2)


def load_inventory(path) -> List[StockSheet]:
    """Read a stock list from `path`, Excel or (legacy) JSON, decided by the
    file extension."""
    if _is_excel_path(path):
        return load_inventory_xlsx(path)
    return load_inventory_json(path)


def save_inventory(path, stock: List[StockSheet]):
    """Write `stock` back to `path` in whatever format that path names --
    so a job committed against a legacy .json inventory stays .json instead
    of quietly becoming a workbook the caller never asked for."""
    if _is_excel_path(path):
        save_inventory_xlsx(path, stock)
    else:
        save_inventory_json(path, stock)


def convert_inventory(src, dest=None) -> str:
    """Rewrite an inventory in the other format -- in practice, an old
    .json one as the .xlsx everything now expects. Returns the path written
    (`src` with a .xlsx extension when `dest` is omitted)."""
    if dest is None:
        dest = os.path.splitext(str(src))[0] + DEFAULT_INVENTORY_EXTENSION
    save_inventory(dest, load_inventory(src))
    return dest


GroupKey = Tuple[str, float]


def group_parts(parts: List[Part]) -> Dict[GroupKey, List[Part]]:
    """Group by (material, thickness). Parts missing either fall into
    ("unspecified", 0.0) -- callers should fill those in before nesting for
    real, since that bucket is a signal the source data is incomplete, not
    a real material."""
    groups: Dict[GroupKey, List[Part]] = {}
    for p in parts:
        key = (p.material or "unspecified", p.thickness if p.thickness is not None else 0.0)
        groups.setdefault(key, []).append(p)
    return groups


def _matches(stock: StockSheet, material: str, thickness: float, tol: float = 1e-3) -> bool:
    return stock.material == material and abs(stock.thickness - thickness) <= tol


def candidate_stocks_for(material: str, thickness: float, inventory: List[StockSheet],
                          min_w: float = 0.0, min_h: float = 0.0) -> List[StockSheet]:
    """Every stock entry matching (material, thickness), with quantity on
    hand and big enough for the group's largest part (either orientation,
    since the part may end up rotated) -- unranked, unlike best_stock_for()
    which already picks a single winner by a static heuristic. Used by
    joint_stock_optimization (and stock_solver.py's deeper search) to
    actually nest against every viable size instead of trusting that
    heuristic's ranking blindly. Not module-private -- stock_solver.py
    calls this directly."""
    return [
        s for s in inventory
        if _matches(s, material, thickness)
        and (s.quantity is None or s.quantity > 0)
        and ((s.width >= min_w and s.height >= min_h) or (s.width >= min_h and s.height >= min_w))
    ]


def best_stock_for(material: str, thickness: float, inventory: List[StockSheet],
                    min_w: float = 0.0, min_h: float = 0.0,
                    prefer_remnants: bool = True) -> Optional[StockSheet]:
    """Remnants first (use up scrap before cutting into new material) when
    `prefer_remnants` is set (the default), then smallest area that's still
    big enough for the group's largest part (checked both orientations,
    since the part may end up rotated). With `prefer_remnants=False`,
    remnant status is ignored entirely and this is pure smallest-fit."""
    candidates = candidate_stocks_for(material, thickness, inventory, min_w, min_h)
    if not candidates:
        return None
    candidates.sort(key=lambda s: (prefer_remnants and not s.is_remnant, s.area()))
    return candidates[0]


def nest_against_stock(stock: StockSheet, remaining: List[Part], kerf: float,
                        use_ga: bool, ga_kwargs: Optional[dict], nester_kwargs: dict,
                        should_stop=None):
    """Nests `remaining` against one candidate stock size, truncating to
    on-hand quantity (overflow sheets' parts fall back to unplaced, same
    truncation the main cascade loop always applied inline before this was
    pulled out into its own function). Returns (sheets, unplaced_names,
    note_or_None). Not module-private -- stock_solver.py builds its
    "fill exactly one sheet" primitive on top of this."""
    if use_ga:
        sheets, unplaced_names = optimize_order(
            stock.width, stock.height, remaining, kerf=kerf, should_stop=should_stop,
            **(ga_kwargs or {}), **nester_kwargs
        )
    else:
        n = Nester(stock.width, stock.height, kerf=kerf, **nester_kwargs)
        for p in remaining:
            n.add_part(p)
        sheets, unplaced_names = n.run()
    sheets = [s for s in sheets if s]
    unplaced_names = list(unplaced_names)

    note = None
    if stock.quantity is not None and len(sheets) > stock.quantity:
        # Geometrically these parts fit on len(sheets) sheets of this
        # stock, but there isn't physically that much on hand -- the
        # overflow sheets' worth of parts fall back to the next-best stock
        # instead, same as a geometric miss.
        overflow_sheets, sheets = sheets[stock.quantity:], sheets[:stock.quantity]
        for s in overflow_sheets:
            unplaced_names.extend(pp.name for pp in s)
        note = (f"only {stock.quantity} of stock {stock.id or f'{stock.width}x{stock.height}'} "
                f"on hand -- overflow retried on next-best stock")
    return sheets, unplaced_names, note


def _stock_objective(stock: StockSheet, sheets: List[List[PlacedPart]], prefer_remnants: bool = True):
    """Ranking key for joint_stock_optimization. When `prefer_remnants` is
    set (the default), any remnant candidate outranks any non-remnant one
    outright, before sheet count/cost/waste are even compared -- matching
    best_stock_for()'s own always-remnants-first default, so this "smarter,
    measure every candidate" mode can't quietly trade away a usable remnant
    for a full sheet that merely scores a little better on cost/waste.
    Below that (or always, if `prefer_remnants=False`): fewer sheets wins;
    among ties, lower total cost if `material_cost()` is computable
    (unpriced = +inf, so an unpriced candidate never wins a tie against a
    priced one, but a comparison between two unpriced candidates still
    falls through to wasted area cleanly); among ties on both, less wasted
    area."""
    n_sheets = len(sheets)
    per_sheet_cost = stock.material_cost()
    cost = n_sheets * per_sheet_cost if per_sheet_cost is not None else float("inf")
    used_area = sum(pp.net_area() for s in sheets for pp in s)
    wasted = stock.area() * max(n_sheets, 1) - used_area
    remnant_rank = 0 if (prefer_remnants and stock.is_remnant) else 1
    return (remnant_rank, n_sheets, cost, wasted)


@dataclass
class JobResult:
    material: str
    thickness: float
    sheets: List[List[PlacedPart]] = field(default_factory=list)
    sheet_stock: List[StockSheet] = field(default_factory=list)  # sheet_stock[i] is the stock sheets[i] was cut from
    unplaced: List[str] = field(default_factory=list)
    notes: List[str] = field(default_factory=list)


def names_to_parts(names: List[str], template_by_name: Dict[str, Part]) -> List[Part]:
    """Rebuild Part objects (with quantity = occurrence count) from a flat
    list of part names -- what Nester.run()'s `unplaced` return gives back
    -- so they can be retried against a different stock size. Not
    module-private -- stock_solver.py reuses this to carry a plan slot's
    leftover parts forward to the next slot."""
    counts: Dict[str, int] = {}
    for name in names:
        counts[name] = counts.get(name, 0) + 1
    rebuilt = []
    for name, qty in counts.items():
        template = template_by_name[name]
        rebuilt.append(Part(
            name=template.name, points=template.points, holes=template.holes,
            quantity=qty, rotations=template.rotations,
            material=template.material, thickness=template.thickness,
        ))
    return rebuilt


def group_bbox(parts: List[Part]) -> Tuple[float, float]:
    """Largest width/height (independently) across `parts`' own bounding
    boxes -- the minimum a stock sheet must offer (in some orientation) to
    have any chance of fitting the group's biggest part. Extracted out of
    run_job()'s cascade loop so stock_solver.py's plan search can compute
    the same thing without duplicating it."""
    w = max(max(x for x, _ in p.points) - min(x for x, _ in p.points) for p in parts)
    h = max(max(y for _, y in p.points) - min(y for _, y in p.points) for p in parts)
    return w, h


def greedy_best_stock_step(material: str, thickness: float, remaining: List[Part],
                            inventory: List[StockSheet], kerf: float, use_ga: bool,
                            ga_kwargs: Optional[dict], nester_kwargs: dict, should_stop,
                            min_w: float, min_h: float, prefer_remnants: bool = True):
    """One "which single stock size is actually best for what's left"
    decision -- nests `remaining` against EVERY viable candidate size (not
    just whichever `best_stock_for()`'s static heuristic picks first) and
    keeps whichever yields the best `_stock_objective()` (remnants first if
    `prefer_remnants`, then fewest sheets, then lowest cost if priced, then
    least waste). This is the `joint_stock_optimization` algorithm itself;
    pulled out of run_job()'s cascade loop so stock_solver.py's deeper plan
    search can reuse the exact same per-step decision as its construction
    baseline. Returns (stock_or_None, sheets, unplaced_names, note_or_None)."""
    candidates = candidate_stocks_for(material, thickness, inventory, min_w, min_h)
    best = None
    for cand in candidates:
        if should_stop is not None and should_stop():
            break
        cand_sheets, cand_unplaced, cand_note = nest_against_stock(
            cand, remaining, kerf, use_ga, ga_kwargs, nester_kwargs, should_stop)
        obj = _stock_objective(cand, cand_sheets, prefer_remnants)
        if best is None or obj < best[0]:
            best = (obj, cand, cand_sheets, cand_unplaced, cand_note)
    if best is None:
        return None, None, None, None
    _obj, stock, sheets, unplaced_names, note = best
    return stock, sheets, unplaced_names, note


def run_job(parts: List[Part], inventory: List[StockSheet], kerf: float = 0.0,
            max_stock_switches: int = 20, use_ga: bool = False, ga_kwargs: Optional[dict] = None,
            joint_stock_optimization: bool = False, true_joint_stock_optimization: bool = False,
            prefer_remnants: bool = True, should_stop=None, **nester_kwargs) -> List[JobResult]:
    """One JobResult per (material, thickness) group found in `parts`.
    Within a group, cascades from best-matching stock to the next as each
    one is exhausted (geometrically or by on-hand quantity) -- see module
    docstring.

    `prefer_remnants` (default True) makes remnants always outrank full
    sheets when choosing stock, REGARDLESS of which of the three modes
    below picks the candidate -- `best_stock_for()`'s default heuristic
    already did this; without threading it through here too,
    `joint_stock_optimization`/`true_joint_stock_optimization` could
    silently override it by picking a full sheet that merely scores a
    little better on sheets/cost/waste than a perfectly usable remnant.
    Set False to let those two modes rank candidates purely on
    sheets/cost/waste, ignoring remnant status entirely.

    If `use_ga` is set, each pass against a stock size runs
    `genetic.optimize_order()` instead of plain `Nester.run()` -- i.e. one
    GA search per (material, thickness) group *per stock size it cascades
    through*, not just once per group. That's the right place for it: a
    part ordering that's great for one stock size can be a poor one for
    the next-smaller size the leftovers cascade to, so re-searching per
    pass beats searching once up front. `ga_kwargs` is forwarded to
    `optimize_order()` (population_size, generations, time_budget_seconds,
    etc.) -- see its docstring for the slow/thorough trade-off this
    implies versus the default plain largest-first ordering.

    If `joint_stock_optimization` is set, each cascade pass nests the
    remaining parts against EVERY viable stock size for the group (not
    just whichever `best_stock_for()`'s remnants-first/smallest-first
    heuristic picks first) and keeps whichever candidate actually yields
    the best result -- fewest sheets, then lowest total material cost if
    `price_per_kg`/`density_g_cm3` are set (see `StockSheet.material_cost()`
    and `_stock_objective()`). HONEST COST: this multiplies the number of full
    nesting passes (or full GA searches, if `use_ga` is also set) by
    however many candidate stock sizes exist for the group, per cascade
    step -- meaningfully slower than the default heuristic, which only
    ever nests against the one size it already committed to. HONEST
    LIMITATION: still not a true joint solve across MIXED sizes within one
    pass (it still can't decide "2 sheets of A + 1 of B" beats "3 of B" in
    one shot) -- it upgrades *which single size* each cascade pass commits
    to, from a static heuristic guess to an actually-measured best, not a
    full cutting-stock solver.

    If `true_joint_stock_optimization` is set, each group is handed to
    `stock_solver.optimize_stock_plan()` ONCE instead of running this
    function's own per-step cascade at all: that module hill-climbs over
    an explicit, ordered PLAN of specific stock sheets (possibly several
    different sizes) for the group's WHOLE remaining part list at once,
    seeded from exactly this function's `joint_stock_optimization`
    algorithm (via the new `greedy_best_stock_step()` helper both now
    share) so it can never do worse -- see `stock_solver.py`'s module
    docstring for the search itself and its honest cost. This is the
    actual "2 sheets of A + 1 of B vs. 3 of B" joint decision the plain
    `joint_stock_optimization` flag above still can't make; when this flag
    is set, `joint_stock_optimization`'s value is ignored (the deeper
    search always uses that algorithm as its own starting point)."""
    results = []
    for (material, thickness), group in group_parts(parts).items():
        if should_stop is not None and should_stop():
            break
        result = JobResult(material=material, thickness=thickness)

        if true_joint_stock_optimization:
            import stock_solver  # local: stock_solver imports this module too
            sheets, sheet_stock, unplaced_names, notes = stock_solver.optimize_stock_plan(
                material, thickness, group, inventory, kerf=kerf, use_ga=use_ga,
                ga_kwargs=ga_kwargs, nester_kwargs=nester_kwargs,
                max_stock_switches=max_stock_switches, prefer_remnants=prefer_remnants,
                should_stop=should_stop,
            )
            result.sheets, result.sheet_stock = sheets, sheet_stock
            result.unplaced, result.notes = unplaced_names, notes
            counts: Dict[int, int] = {}
            for s in sheet_stock:
                counts[id(s)] = counts.get(id(s), 0) + 1
            decremented = set()
            for s in sheet_stock:
                if s.quantity is not None and id(s) not in decremented:
                    decremented.add(id(s))
                    s.quantity = max(0, s.quantity - counts[id(s)])
            results.append(result)
            continue

        template_by_name = {p.name: p for p in group}
        remaining = list(group)

        for _ in range(max_stock_switches):
            if should_stop is not None and should_stop():
                result.notes.append("nesting stopped by user; remaining parts were not evaluated")
                break
            if not remaining:
                break
            min_w, min_h = group_bbox(remaining)

            if joint_stock_optimization:
                stock, sheets, unplaced_names, note = greedy_best_stock_step(
                    material, thickness, remaining, inventory, kerf, use_ga,
                    ga_kwargs, nester_kwargs, should_stop, min_w, min_h, prefer_remnants)
            else:
                stock = best_stock_for(material, thickness, inventory, min_w, min_h, prefer_remnants)
                sheets, unplaced_names, note = (
                    nest_against_stock(stock, remaining, kerf, use_ga, ga_kwargs, nester_kwargs, should_stop)
                    if stock is not None else (None, None, None)
                )

            if stock is None:
                result.unplaced.extend(p.name for p in remaining for _ in range(p.quantity))
                result.notes.append("no (more) matching stock in inventory for this material/thickness")
                remaining = []
                break
            if note:
                result.notes.append(note)

            result.sheets.extend(sheets)
            result.sheet_stock.extend([stock] * len(sheets))
            if stock.quantity is not None:
                stock.quantity = max(0, stock.quantity - len(sheets))

            if not unplaced_names:
                remaining = []
                break
            if len(unplaced_names) == sum(p.quantity for p in remaining):
                # No progress at all this round (stock too small / no
                # sheets fit even one part) -- retrying the same inventory
                # would just loop forever between the same two outcomes.
                result.unplaced.extend(unplaced_names)
                result.notes.append("remaining parts don't fit any matching stock (including this one)")
                remaining = []
                break
            remaining = names_to_parts(unplaced_names, template_by_name)
        else:
            result.unplaced.extend(p.name for p in remaining for _ in range(p.quantity))
            result.notes.append(f"gave up after {max_stock_switches} stock switches")

        results.append(result)
    return results


def _stock_key(s: StockSheet) -> str:
    # Two stock entries sharing (material, thickness, width, height) with
    # no id are treated as the same SKU for logging purposes -- if they
    # were meant to be tracked separately they'd have their own ids.
    return s.id or f"{s.material}|{s.thickness}|{s.width}x{s.height}"


def commit_job(parts: List[Part], inventory_path: str, kerf: float = 0.0,
               log_path: Optional[str] = None, use_ga: bool = False,
               ga_kwargs: Optional[dict] = None, joint_stock_optimization: bool = False,
               true_joint_stock_optimization: bool = False, prefer_remnants: bool = True,
               capture_remnants: bool = True,
               min_remnant_dimension: float = 100.0, **nester_kwargs) -> List[JobResult]:
    """Like run_job(), but for real: loads inventory fresh from
    `inventory_path`, runs the job against those actual StockSheet objects
    (mutating their `.quantity` in place, same as run_job() always did),
    persists the updated counts back to that same file, and appends one
    line per stock consumed -- plus one per group with leftover unplaced
    parts -- to a JSON-lines audit log (default: `inventory_path` with
    '.log.jsonl' appended).

    This is the one function in this module that actually changes what's
    on disk -- run_job() alone (e.g. via the workbench UI's plain "Run
    Nesting") only mutates an in-memory list, safe to call repeatedly while
    experimenting. Call this only once you mean to actually commit to
    having cut those sheets.

    If `capture_remnants` is set (the default), each cut sheet's leftover
    offcut is recorded back into `stock` as a new remnant -- see
    remnant.py's module docstring for why that's the largest axis-aligned
    empty rectangle among placed parts' bounding boxes, not their exact
    (possibly concave) outlines: this project's stock model is rectangles
    only, so an exact polygon remnant would have nothing to be stored as.
    A remnant smaller than `min_remnant_dimension` on either side is
    dropped rather than recorded -- not worth tracking as reusable stock."""
    stock = load_inventory(inventory_path)
    results = run_job(parts, stock, kerf=kerf, use_ga=use_ga, ga_kwargs=ga_kwargs,
                       joint_stock_optimization=joint_stock_optimization,
                       true_joint_stock_optimization=true_joint_stock_optimization,
                       prefer_remnants=prefer_remnants, **nester_kwargs)

    new_remnants: List[StockSheet] = []
    if capture_remnants:
        for r in results:
            for sheet, s in zip(r.sheets, r.sheet_stock):
                if not sheet:
                    continue
                obstacle_bboxes = [polygon_bbox(pp.points) for pp in sheet]
                rect = largest_empty_rect(s.width, s.height, obstacle_bboxes)
                if rect is None:
                    continue
                _x, _y, rem_w, rem_h = rect
                if rem_w < min_remnant_dimension or rem_h < min_remnant_dimension:
                    continue
                remnant = StockSheet(
                    material=r.material, thickness=r.thickness,
                    width=rem_w, height=rem_h, quantity=1, is_remnant=True,
                    id=f"REM-{r.material}-{uuid.uuid4().hex[:8]}",
                    # Same material as the sheet it was cut from, so the
                    # same rates apply -- lets a later job value "money
                    # saved" by using this remnant instead of buying new,
                    # and value ITS OWN eventual offcuts as scrap too (see
                    # nesting_widgets.py's job-financials calculation).
                    price_per_kg=s.price_per_kg, density_g_cm3=s.density_g_cm3,
                    scrap_price_per_kg=s.scrap_price_per_kg,
                )
                stock.append(remnant)
                new_remnants.append(remnant)

    save_inventory(inventory_path, stock)

    if log_path is None:
        log_path = inventory_path + ".log.jsonl"
    now = datetime.datetime.now(datetime.timezone.utc).isoformat()
    with open(log_path, "a", encoding="utf-8") as f:
        for r in results:
            counts: Dict[str, List] = {}
            for s in r.sheet_stock:
                key = _stock_key(s)
                if key not in counts:
                    counts[key] = [0, s]
                counts[key][0] += 1
            for n_sheets, s in counts.values():
                f.write(json.dumps({
                    "timestamp": now, "event": "consumed",
                    "material": r.material, "thickness": r.thickness,
                    "stock_id": s.id, "stock_size": f"{s.width}x{s.height}",
                    "is_remnant": s.is_remnant,
                    "sheets_used": n_sheets, "remaining_qty": s.quantity,
                }) + "\n")
            if r.unplaced:
                f.write(json.dumps({
                    "timestamp": now, "event": "unplaced",
                    "material": r.material, "thickness": r.thickness,
                    "parts": r.unplaced, "notes": r.notes,
                }) + "\n")
        for remnant in new_remnants:
            f.write(json.dumps({
                "timestamp": now, "event": "remnant_created",
                "material": remnant.material, "thickness": remnant.thickness,
                "stock_id": remnant.id, "stock_size": f"{remnant.width}x{remnant.height}",
            }) + "\n")
    return results
