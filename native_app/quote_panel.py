"""
quote_panel.py
--------------
The native app's Quote tab (docs/QUOTING_PLAN.md, Phase 4 -- the UI part):
turns the job's parts into a priced quote with the `quoting/` engine.

  * Lines -- every part in the Parts tab becomes a sheet line (kept in step
    as parts are added or removed; what you type here per line survives
    that). Bought-in parts, tube/section members and a welded assembly
    are added from the ribbon. The form under the list edits whichever
    line is selected: the facts the geometry can't give (thickness of a
    DXF part, bend count, tapped holes, paint, profile and length of a
    tube, unit cost of a bolt, weld length of an assembly).
  * Settings -- quantity breaks, material from a real nest or a quick
    utilisation estimate, allocation, markup per cost area, paint and
    transport. Rates start as the engine's "calibrate me" starter table;
    Load Rates reads a shop's own (.xlsx or SQLite), Export Rates writes
    the current table to .xlsx to edit.
  * Results -- one row per quantity break (cost, price, price per set),
    and for the selected break the full breakdown: a column per cost
    centre, per-unit, every cell editable. A typed value overrides the
    cell for that break (shown bold, the calculated value in its
    tooltip); clearing it goes back to the calculated one.

Sheet material comes from the stock loaded on the Stock tab. With no stock
loaded, each material/thickness gets a stand-in 3000x1500 sheet at the
"Sheet price/kg" setting, and the quote says so.

Pricing runs on a worker thread -- re-nesting every break of a big job
takes a while -- and the tab stays usable meanwhile. Everything here reads
the panel's parts and stock; nothing is written back to the job.
"""

import copy
import math
import os
from dataclasses import dataclass, field
from typing import Dict, List, Optional

from PySide6 import QtCore, QtGui, QtWidgets

from bar_nest import BarStock
from inventory import StockSheet
from quoting.costing import MATERIAL, QUOTE, CostingSettings, QuoteEngine
from quoting.features import (Bend, HoleFeature, PartFeatures, ProfileInfo, default_density,
                              features_from_polygons)
from quoting.operations import default_operations, painting
from quoting.ops_assembly import QuoteCost
from quoting.rates import load_rates, save_rates, starter_rates

LINE_TYPES = ("sheet", "tube", "purchased", "assembly")
TYPE_LABELS = {"sheet": "Sheet", "tube": "Tube / section", "section": "Tube / section",
               "purchased": "Bought-in", "assembly": "Assembly"}
PROFILE_KINDS = ("RHS", "SHS", "CHS", "FLAT", "ANGLE")
END_CUTS = ("square", "mitre", "cope")
TAP_DRILL = {"M3": 2.5, "M4": 3.3, "M5": 4.2, "M6": 5.0, "M8": 6.8, "M10": 8.5, "M12": 10.2}
COST_AREAS = ("material", "cutting", "bending", "machining", "welding", "assembly",
              "finishing", "purchased", "quote")
DEFAULT_MARKUPS = {"material": 10.0, "purchased": 25.0}
DEFAULT_MARKUP = 25.0
FALLBACK_SHEET = (3000.0, 1500.0)


# ------------------------------------------------------------------- model

@dataclass
class LineSpec:
    """One quote line as the user edits it. `to_features()` turns it into
    the engine's PartFeatures."""
    name: str
    part_type: str = "sheet"
    quantity: int = 1
    material: str = ""
    thickness: Optional[float] = None
    from_parts: bool = False              # mirrors a row of the Parts tab
    # sheet
    outer: list = field(default_factory=list)
    holes: list = field(default_factory=list)
    bends: Optional[int] = None           # None = unknown
    bend_details: list = field(default_factory=list)
    tapped_holes: int = 0
    tap_size: str = "M8"
    paint: bool = False
    # tube / section
    profile_kind: str = "RHS"
    width: float = 100.0
    height: float = 50.0
    wall: float = 3.0
    length: float = 1000.0
    end_a: str = "square"
    end_b: str = "square"
    bar_length: float = 6000.0
    price_per_m: float = 10.0
    # bought-in
    unit_cost: float = 0.0
    supplier: str = ""
    markup_pct: Optional[float] = None
    # assembly
    part_count: int = 0
    fastener_count: int = 0
    weld_length: float = 0.0
    weld_type: str = "fillet"

    # -- tube geometry -------------------------------------------------
    def profile(self) -> ProfileInfo:
        w, h, t = self.width, self.height, self.wall
        kind = self.profile_kind
        if kind == "SHS":
            h = w
        if kind == "CHS":
            area = math.pi / 4 * (w ** 2 - max(w - 2 * t, 0) ** 2)
            return ProfileInfo("CHS", w, w, t, area, math.pi * w)
        if kind in ("RHS", "SHS"):
            area = w * h - max(w - 2 * t, 0) * max(h - 2 * t, 0)
            return ProfileInfo(kind, max(w, h), min(w, h), t, area, 2 * (w + h))
        if kind == "FLAT":
            return ProfileInfo("FLAT", max(w, t), min(w, t), t, w * t, 2 * (w + t))
        return ProfileInfo("ANGLE", max(w, h), min(w, h), t, t * (w + h - t), 2 * (w + h))

    def to_features(self) -> PartFeatures:
        material = self.material.strip() or None
        if self.part_type == "sheet":
            if self.bend_details:
                bends = [Bend(d.get("angle_deg"), d.get("length"), d.get("radius"), d.get("direction"))
                         for d in self.bend_details]
                if self.bends is not None and self.bends != len(bends):
                    bends = [Bend() for _ in range(self.bends)]
            else:
                bends = None if self.bends is None else [Bend() for _ in range(self.bends)]
            holes = [HoleFeature(TAP_DRILL.get(self.tap_size, 6.8), "tapped", self.tap_size,
                                 count=self.tapped_holes)] if self.tapped_holes else []
            f = features_from_polygons(self.name, self.outer, self.holes, material=material,
                                       thickness=self.thickness, quantity=self.quantity,
                                       bends=bends, hole_features=holes)
            f.extra["paint"] = self.paint
            return f
        if self.part_type == "tube":
            prof = self.profile()
            f = PartFeatures(self.name, part_type="section" if not prof.hollow else "tube",
                             material=material, quantity=self.quantity, profile=prof,
                             length=self.length, end_cuts=(self.end_a, self.end_b),
                             thickness=self.wall)
            f.extra["paint"] = self.paint
            return f
        if self.part_type == "purchased":
            return PartFeatures(self.name, part_type="purchased", quantity=self.quantity,
                                unit_cost=self.unit_cost, supplier=self.supplier or None,
                                markup=None if self.markup_pct is None else self.markup_pct / 100.0)
        return PartFeatures(self.name, part_type="assembly", quantity=self.quantity, material=material,
                            thickness=self.thickness, part_count=self.part_count,
                            fastener_count=self.fastener_count,
                            manual_weld_length=self.weld_length or None)

    def bar_stock(self) -> Optional[BarStock]:
        if self.part_type != "tube":
            return None
        prof = self.profile()
        density = default_density(self.material)
        kg_per_m = prof.area * density / 1000.0 if density and prof.area else None
        return BarStock(self.material or "", prof.designation, self.bar_length,
                        price_per_m=self.price_per_m, kg_per_m=kg_per_m)

    def summary(self) -> str:
        if self.part_type == "sheet":
            t = f"{self.thickness:g} mm" if self.thickness else "? mm"
            bends = "bends ?" if self.bends is None else f"{self.bends} bend(s)"
            extra = [f"{self.tapped_holes}x {self.tap_size}"] if self.tapped_holes else []
            if self.paint:
                extra.append("painted")
            return ", ".join([t, bends] + extra)
        if self.part_type == "tube":
            return f"{self.profile().designation} x {self.length:g} mm, ends {self.end_a}/{self.end_b}"
        if self.part_type == "purchased":
            return f"{self.unit_cost:.2f} each" + (f" from {self.supplier}" if self.supplier else "")
        return (f"{self.part_count} parts, {self.fastener_count} fasteners, "
                f"{self.weld_length:g} mm weld")


@dataclass
class QuoteSetup:
    breaks: List[int] = field(default_factory=lambda: [1, 10, 100])
    material_mode: str = "nest"
    allocation: str = "net_area"
    tube_machine: str = "saw"
    markups: Dict[str, float] = field(default_factory=lambda: dict(DEFAULT_MARKUPS))  # percent
    default_markup: float = DEFAULT_MARKUP
    paint_rate: float = 15.0
    paint_setup: float = 25.0
    transport: float = 80.0
    fallback_price_per_kg: float = 1.20
    kerf: float = 0.0


def parse_breaks(text: str) -> List[int]:
    """"1, 10, 100" -> [1, 10, 100]; raises ValueError on anything else."""
    out = []
    for tok in text.replace(";", ",").replace(" ", ",").split(","):
        if tok.strip():
            v = int(tok)
            if v < 1:
                raise ValueError("quantities must be 1 or more")
            out.append(v)
    if not out:
        raise ValueError("enter at least one quantity")
    return sorted(set(out))


def build_engine(lines: List[LineSpec], setup: QuoteSetup, rates, stock: Optional[List[StockSheet]]):
    """(engine, features, notes) for `lines` -- Qt-free so it can be tested
    and run on a worker thread."""
    notes = []
    features = [l.to_features() for l in lines]
    sheet_stock = list(stock or [])
    if not sheet_stock:
        seen = set()
        for f in features:
            if f.part_type == "sheet" and f.material and f.thickness:
                key = (f.material, f.thickness)
                if key in seen:
                    continue
                seen.add(key)
                sheet_stock.append(StockSheet(f.material, f.thickness, *FALLBACK_SHEET,
                                              price_per_kg=setup.fallback_price_per_kg,
                                              density_g_cm3=f.density()))
        if seen:
            notes.append(f"No stock loaded on the Stock tab: sheet material is priced on "
                         f"{FALLBACK_SHEET[0]:g}x{FALLBACK_SHEET[1]:g} sheets at "
                         f"{setup.fallback_price_per_kg:g}/kg.")
    bars, seen_bars = [], set()
    for l in lines:
        b = l.bar_stock()
        if b is not None and (b.material, b.profile) not in seen_bars:
            seen_bars.add((b.material, b.profile))
            bars.append(b)
    ops = default_operations(rates, tube_machine=setup.tube_machine) + [
        painting(setup.paint_rate, setup=setup.paint_setup)]
    settings = CostingSettings(
        allocation=setup.allocation, material_mode=setup.material_mode, kerf=setup.kerf,
        markups={k: v / 100.0 for k, v in setup.markups.items()},
        default_markup=setup.default_markup / 100.0)
    quote_costs = [QuoteCost("transport", setup.transport)] if setup.transport else []
    return QuoteEngine(rates, sheet_stock, bars, operations=ops, settings=settings,
                       quote_costs=quote_costs), features, notes


# ------------------------------------------------------------------ worker

class _PriceWorker(QtCore.QThread):
    priced = QtCore.Signal(object, object)   # (QuoteResult, extra notes)
    failed = QtCore.Signal(str)

    def __init__(self, lines, setup, rates, stock, parent=None):
        super().__init__(parent)
        self._args = (lines, setup, rates, stock)

    def run(self):
        lines, setup, rates, stock = self._args
        try:
            engine, features, notes = build_engine(lines, setup, rates, stock)
            result = engine.price(features, breaks=setup.breaks)
        except Exception as e:  # noqa: BLE001 -- shown to the user, never swallowed
            self.failed.emit(str(e))
            return
        self.priced.emit(result, notes)


# ------------------------------------------------------------------- panel

def _spin(lo, hi, decimals=2, step=1.0, suffix=""):
    s = QtWidgets.QDoubleSpinBox()
    s.setRange(lo, hi)
    s.setDecimals(decimals)
    s.setSingleStep(step)
    if suffix:
        s.setSuffix(suffix)
    s.setKeyboardTracking(False)
    return s


def _ispin(lo, hi):
    s = QtWidgets.QSpinBox()
    s.setRange(lo, hi)
    s.setKeyboardTracking(False)
    return s


class QuotePanel(QtWidgets.QWidget):
    busy_changed = QtCore.Signal(bool)
    quote_ready = QtCore.Signal()

    LINE_COLS = ("Line", "Type", "Qty / set", "Material", "Details")

    def __init__(self, panel, parent=None):
        super().__init__(parent)
        self.setObjectName("QuotePanel")
        self.panel = panel
        self.lines: List[LineSpec] = []
        self.setup = QuoteSetup()
        self.rates = starter_rates()
        self.rates_source = "Starter rates (calibrate me)"
        self.result = None
        self.result_notes: List[str] = []
        self.overrides: Dict[tuple, float] = {}
        self._worker = None
        self._loading = False
        self._filling_grid = False

        self._build()
        panel.parts_changed.connect(self.sync_from_parts)
        panel.currency_changed.connect(lambda *_: self._fill_results())
        self.sync_from_parts()

    # ------------------------------------------------------------ layout

    def _build(self):
        outer = QtWidgets.QVBoxLayout(self)
        outer.setContentsMargins(8, 8, 8, 8)
        split = QtWidgets.QSplitter(QtCore.Qt.Vertical)
        split.setChildrenCollapsible(False)
        outer.addWidget(split)

        top = QtWidgets.QSplitter(QtCore.Qt.Horizontal)
        top.setChildrenCollapsible(False)
        split.addWidget(top)

        # lines + form
        left = QtWidgets.QWidget()
        lv = QtWidgets.QVBoxLayout(left)
        lv.setContentsMargins(0, 0, 0, 0)
        lv.addWidget(self._heading("Quote lines"))
        self.lines_table = QtWidgets.QTableWidget(0, len(self.LINE_COLS))
        self.lines_table.setHorizontalHeaderLabels(self.LINE_COLS)
        self.lines_table.setSelectionBehavior(QtWidgets.QAbstractItemView.SelectRows)
        self.lines_table.setSelectionMode(QtWidgets.QAbstractItemView.SingleSelection)
        self.lines_table.setEditTriggers(QtWidgets.QAbstractItemView.NoEditTriggers)
        self.lines_table.verticalHeader().hide()
        self.lines_table.horizontalHeader().setStretchLastSection(True)
        self.lines_table.itemSelectionChanged.connect(self._on_line_selected)
        self.lines_table.setMinimumHeight(150)
        lv.addWidget(self.lines_table, 3)
        self.form_box = QtWidgets.QGroupBox("Selected line")
        lv.addWidget(self.form_box, 2)
        self._build_form()
        top.addWidget(left)

        top.addWidget(self._build_settings())
        top.setStretchFactor(0, 3)
        top.setStretchFactor(1, 1)

        split.addWidget(self._build_results())
        split.setStretchFactor(0, 1)
        split.setStretchFactor(1, 1)
        self.split = split

    @staticmethod
    def _heading(text):
        lab = QtWidgets.QLabel(text)
        f = lab.font()
        f.setBold(True)
        lab.setFont(f)
        return lab

    def _build_form(self):
        box = QtWidgets.QVBoxLayout(self.form_box)
        common = QtWidgets.QFormLayout()
        common.setFieldGrowthPolicy(QtWidgets.QFormLayout.AllNonFixedFieldsGrow)
        row = QtWidgets.QHBoxLayout()
        self.f_name = QtWidgets.QLineEdit()
        self.f_qty = _ispin(1, 100000)
        self.f_material = QtWidgets.QLineEdit()
        self.f_material.setPlaceholderText("e.g. mild steel")
        self.f_thickness = _spin(0, 200, 2, 0.5, " mm")
        self.f_thickness.setSpecialValueText("?")
        for lab, w in (("Name", self.f_name), ("Qty / set", self.f_qty), ("Material", self.f_material),
                       ("Thickness", self.f_thickness)):
            row.addWidget(QtWidgets.QLabel(lab))
            row.addWidget(w, 2 if w in (self.f_name, self.f_material) else 1)
        common.addRow(row)
        box.addLayout(common)

        self.form_pages = QtWidgets.QStackedWidget()
        box.addWidget(self.form_pages)
        self.form_placeholder = QtWidgets.QLabel("Select a line to edit it.")
        self.form_placeholder.setAlignment(QtCore.Qt.AlignCenter)
        self.form_pages.addWidget(self.form_placeholder)

        # sheet
        page = QtWidgets.QWidget()
        g = QtWidgets.QGridLayout(page)
        self.f_bends = _ispin(-1, 99)
        self.f_bends.setSpecialValueText("unknown")
        self.f_tapped = _ispin(0, 999)
        self.f_tap_size = QtWidgets.QComboBox()
        self.f_tap_size.addItems(list(TAP_DRILL))
        self.f_paint = QtWidgets.QCheckBox("Painted (both sides)")
        self.f_geom = QtWidgets.QLabel()
        self.f_geom.setWordWrap(True)
        g.addWidget(QtWidgets.QLabel("Bends"), 0, 0)
        g.addWidget(self.f_bends, 0, 1)
        g.addWidget(QtWidgets.QLabel("Tapped holes"), 0, 2)
        g.addWidget(self.f_tapped, 0, 3)
        g.addWidget(self.f_tap_size, 0, 4)
        g.addWidget(self.f_paint, 0, 5)
        g.addWidget(self.f_geom, 1, 0, 1, 6)
        self.page_sheet = page
        self.form_pages.addWidget(page)

        # tube
        page = QtWidgets.QWidget()
        g = QtWidgets.QGridLayout(page)
        self.f_kind = QtWidgets.QComboBox()
        self.f_kind.addItems(PROFILE_KINDS)
        self.f_width = _spin(1, 2000, 1, 5, " mm")
        self.f_height = _spin(1, 2000, 1, 5, " mm")
        self.f_wall = _spin(0.5, 100, 1, 0.5, " mm")
        self.f_length = _spin(1, 20000, 0, 10, " mm")
        self.f_end_a = QtWidgets.QComboBox()
        self.f_end_a.addItems(END_CUTS)
        self.f_end_b = QtWidgets.QComboBox()
        self.f_end_b.addItems(END_CUTS)
        self.f_bar = _spin(100, 20000, 0, 100, " mm")
        self.f_price_m = _spin(0, 100000, 2, 1)
        self.f_tube_paint = QtWidgets.QCheckBox("Painted")
        cells = (("Profile", self.f_kind), ("Width / OD", self.f_width), ("Height", self.f_height),
                 ("Wall", self.f_wall), ("Cut length", self.f_length), ("End A", self.f_end_a),
                 ("End B", self.f_end_b), ("Bar length", self.f_bar), ("Price / m", self.f_price_m))
        for i, (lab, w) in enumerate(cells):
            g.addWidget(QtWidgets.QLabel(lab), i // 5, (i % 5) * 2)
            g.addWidget(w, i // 5, (i % 5) * 2 + 1)
        g.addWidget(self.f_tube_paint, 1, 8, 1, 2)
        self.page_tube = page
        self.form_pages.addWidget(page)

        # bought-in
        page = QtWidgets.QWidget()
        g = QtWidgets.QHBoxLayout(page)
        self.f_unit_cost = _spin(0, 1e7, 2, 0.05)
        self.f_supplier = QtWidgets.QLineEdit()
        self.f_own_markup = QtWidgets.QCheckBox("Own markup")
        self.f_markup = _spin(0, 1000, 1, 5, " %")
        for lab, w in (("Unit cost", self.f_unit_cost), ("Supplier", self.f_supplier)):
            g.addWidget(QtWidgets.QLabel(lab))
            g.addWidget(w)
        g.addWidget(self.f_own_markup)
        g.addWidget(self.f_markup)
        self.page_purchased = page
        self.form_pages.addWidget(page)

        # assembly
        page = QtWidgets.QWidget()
        g = QtWidgets.QHBoxLayout(page)
        self.f_parts = _ispin(0, 9999)
        self.f_fasteners = _ispin(0, 9999)
        self.f_weld = _spin(0, 1e6, 0, 50, " mm")
        for lab, w in (("Parts", self.f_parts), ("Fasteners", self.f_fasteners), ("Weld length", self.f_weld)):
            g.addWidget(QtWidgets.QLabel(lab))
            g.addWidget(w)
        g.addStretch(1)
        self.page_assembly = page
        self.form_pages.addWidget(page)

        for w in (self.f_name, self.f_material, self.f_supplier):
            w.editingFinished.connect(self._form_changed)
        for w in (self.f_qty, self.f_thickness, self.f_bends, self.f_tapped, self.f_width, self.f_height,
                  self.f_wall, self.f_length, self.f_bar, self.f_price_m, self.f_unit_cost, self.f_markup,
                  self.f_parts, self.f_fasteners, self.f_weld):
            w.valueChanged.connect(self._form_changed)
        for w in (self.f_tap_size, self.f_kind, self.f_end_a, self.f_end_b):
            w.currentIndexChanged.connect(self._form_changed)
        for w in (self.f_paint, self.f_tube_paint, self.f_own_markup):
            w.toggled.connect(self._form_changed)
        self._set_form_enabled(False)

    def _build_settings(self):
        box = QtWidgets.QGroupBox("Quote settings")
        column = QtWidgets.QVBoxLayout(box)
        form = QtWidgets.QFormLayout()
        column.addLayout(form)
        self.s_breaks = QtWidgets.QLineEdit("1, 10, 100")
        self.s_breaks.setToolTip("Quantity breaks, in sets -- e.g. 1, 10, 100.")
        form.addRow("Quantities", self.s_breaks)
        self.s_mode = QtWidgets.QComboBox()
        self.s_mode.addItem("Real nest", "nest")
        self.s_mode.addItem("Quick estimate", "estimate")
        self.s_mode.setToolTip("Real nest: re-nests the job at every quantity against the loaded stock.\n"
                               "Quick estimate: part area / 75% utilisation -- fast, for big jobs.")
        form.addRow("Material", self.s_mode)
        self.s_alloc = QtWidgets.QComboBox()
        self.s_alloc.addItem("By net area", "net_area")
        self.s_alloc.addItem("By bounding box", "bbox")
        form.addRow("Split sheet cost", self.s_alloc)
        self.s_tube = QtWidgets.QComboBox()
        self.s_tube.addItem("Saw", "saw")
        self.s_tube.addItem("Tube laser", "laser")
        form.addRow("Tube cutting", self.s_tube)
        self.s_price_kg = _spin(0, 1e5, 2, 0.1)
        self.s_price_kg.setValue(self.setup.fallback_price_per_kg)
        self.s_price_kg.setToolTip("Used only when no stock is loaded on the Stock tab.")
        form.addRow("Sheet price / kg", self.s_price_kg)
        self.s_paint = _spin(0, 1e5, 2, 1)
        self.s_paint.setValue(self.setup.paint_rate)
        form.addRow("Paint / m²", self.s_paint)
        self.s_transport = _spin(0, 1e7, 2, 10)
        self.s_transport.setValue(self.setup.transport)
        form.addRow("Transport / quote", self.s_transport)
        for w in (self.s_breaks, self.s_mode, self.s_alloc, self.s_tube, self.s_price_kg, self.s_paint,
                  self.s_transport):
            w.setMinimumWidth(90)
            w.setSizePolicy(QtWidgets.QSizePolicy.Expanding, QtWidgets.QSizePolicy.Fixed)

        self.s_markups = QtWidgets.QTableWidget(len(COST_AREAS), 1)
        self.s_markups.setHorizontalHeaderLabels(["Markup %"])
        self.s_markups.setVerticalHeaderLabels([a.capitalize() for a in COST_AREAS])
        self.s_markups.horizontalHeader().setStretchLastSection(True)
        for i, area in enumerate(COST_AREAS):
            self.s_markups.setItem(i, 0, QtWidgets.QTableWidgetItem(
                f"{self.setup.markups.get(area, self.setup.default_markup):g}"))
        self.s_markups.setToolTip("Markup per cost area -- type a new percentage.")
        column.addWidget(self.s_markups, 1)

        self.rates_label = QtWidgets.QLabel()
        self.rates_label.setWordWrap(True)
        column.addWidget(self.rates_label)
        self._refresh_rates_label()
        scroll = QtWidgets.QScrollArea()
        scroll.setWidgetResizable(True)
        scroll.setFrameShape(QtWidgets.QFrame.NoFrame)
        scroll.setHorizontalScrollBarPolicy(QtCore.Qt.ScrollBarAlwaysOff)
        scroll.setWidget(box)
        scroll.setMinimumWidth(300)
        return scroll

    def _build_results(self):
        w = QtWidgets.QWidget()
        v = QtWidgets.QVBoxLayout(w)
        v.setContentsMargins(0, 0, 0, 0)
        head = QtWidgets.QHBoxLayout()
        head.addWidget(self._heading("Quote"))
        self.status_label = QtWidgets.QLabel("Not priced yet -- click Price Quote on the ribbon.")
        head.addWidget(self.status_label, 1)
        v.addLayout(head)

        self.summary = QtWidgets.QTableWidget(0, 6)
        self.summary.setHorizontalHeaderLabels(["Sets", "Cost", "Price", "Price / set", "Margin", "Status"])
        self.summary.verticalHeader().hide()
        self.summary.setEditTriggers(QtWidgets.QAbstractItemView.NoEditTriggers)
        self.summary.setSelectionBehavior(QtWidgets.QAbstractItemView.SelectRows)
        self.summary.setSelectionMode(QtWidgets.QAbstractItemView.SingleSelection)
        self.summary.horizontalHeader().setSectionResizeMode(QtWidgets.QHeaderView.Stretch)
        self.summary.itemSelectionChanged.connect(self._fill_grid)
        v.addWidget(self.summary)

        self.grid_label = QtWidgets.QLabel()
        v.addWidget(self.grid_label)
        self.grid = QtWidgets.QTableWidget(0, 0)
        self.grid.setMinimumHeight(160)
        self.grid.verticalHeader().hide()
        self.grid.itemChanged.connect(self._on_grid_edit)
        v.addWidget(self.grid, 3)

        self.notes = QtWidgets.QPlainTextEdit()
        self.notes.setReadOnly(True)
        self.notes.setMaximumHeight(70)
        self.notes.setPlaceholderText("Notes about the quote appear here.")
        v.addWidget(self.notes)
        return w

    # ------------------------------------------------------------- lines

    def _app_parts(self):
        """(name, material, qty per assembly, extracted data) per Parts-tab row."""
        panel, out = self.panel, []
        for row in range(panel.table.rowCount()):
            item = panel.table.item(row, 0)
            if item is None:
                continue
            data = panel._extracted.get(item.text())
            if data is None:
                continue
            mat_item = panel.table.item(row, 1)
            material = mat_item.text().strip() if mat_item else ""
            if material.lower() == "unspecified":
                material = ""
            qty_w = panel.table.cellWidget(row, 2)
            out.append((item.text(), material, qty_w.value() if qty_w else 1, data))
        return out

    def sync_from_parts(self):
        """Keep one sheet line per Parts-tab part, carrying over anything
        typed on the Quote tab for a part that's still there."""
        previous = {l.name: l for l in self.lines if l.from_parts}
        synced = []
        for name, material, qty, data in self._app_parts():
            old = previous.get(name)
            spec = LineSpec(name=name, part_type="sheet", quantity=qty, from_parts=True,
                            outer=[tuple(p) for p in data["outer"]],
                            holes=[[tuple(p) for p in h] for h in data["holes"]])
            if old is not None:
                for attr in ("material", "thickness", "bends", "bend_details", "tapped_holes", "tap_size",
                             "paint", "quantity"):
                    setattr(spec, attr, getattr(old, attr))
                if material and not old.material:
                    spec.material = material
            else:
                spec.material = material
                spec.thickness = data.get("thickness")
                bends = data.get("bends")
                spec.bend_details = list(data.get("bend_details") or [])
                spec.bends = (len(spec.bend_details) if spec.bend_details else
                              int(bends) if isinstance(bends, (int, float)) and data.get("method") != "dxf"
                              else None)
                tapped = [h for h in data.get("hole_features") or [] if h.get("kind") == "tapped"]
                if tapped:
                    spec.tapped_holes = sum(int(h.get("count", 1)) for h in tapped)
                    spec.tap_size = tapped[0].get("thread") or spec.tap_size
            synced.append(spec)
        others = [l for l in self.lines if not l.from_parts]
        self.lines = synced + others
        self._fill_lines()

    @staticmethod
    def _material_text(spec):
        if spec.part_type == "purchased":
            return "—"
        return spec.material or "?"

    def _unique_name(self, base):
        names = {l.name for l in self.lines}
        name, n = base, 2
        while name in names:
            name, n = f"{base} {n}", n + 1
        return name

    def add_line(self, part_type):
        defaults = {
            "purchased": dict(name="Bought-in part", unit_cost=0.15),
            "tube": dict(name="Tube", material="mild steel", end_a="mitre", end_b="mitre", length=1200.0),
            "assembly": dict(name="Assembly", material="mild steel", thickness=3.0,
                             part_count=max(1, len(self.lines)), weld_length=400.0),
        }[part_type]
        defaults["name"] = self._unique_name(defaults["name"])
        spec = LineSpec(part_type=part_type, **defaults)
        self.lines.append(spec)
        self._fill_lines()
        self.lines_table.selectRow(len(self.lines) - 1)
        return spec

    def remove_selected_line(self):
        spec = self._selected_line()
        if spec is None:
            return
        if spec.from_parts:
            QtWidgets.QMessageBox.information(
                self, "Quote", f"{spec.name} comes from the Parts tab -- remove it there.")
            return
        self.lines.remove(spec)
        self._fill_lines()

    def _fill_lines(self):
        keep = self._selected_line()
        self.lines_table.setRowCount(len(self.lines))
        for r, l in enumerate(self.lines):
            vals = (l.name, TYPE_LABELS[l.part_type], str(l.quantity), self._material_text(l), l.summary())
            for c, v in enumerate(vals):
                item = QtWidgets.QTableWidgetItem(v)
                if l.from_parts and c == 0:
                    item.setToolTip("From the Parts tab")
                self.lines_table.setItem(r, c, item)
        self.lines_table.resizeColumnsToContents()
        self.lines_table.horizontalHeader().setStretchLastSection(True)
        if keep in self.lines:
            self.lines_table.selectRow(self.lines.index(keep))
        else:
            self._on_line_selected()

    def _selected_line(self) -> Optional[LineSpec]:
        rows = {i.row() for i in self.lines_table.selectedIndexes()} if hasattr(self, "lines_table") else set()
        if len(rows) != 1:
            return None
        r = rows.pop()
        return self.lines[r] if r < len(self.lines) else None

    def _set_form_enabled(self, on):
        for w in (self.f_name, self.f_qty, self.f_material, self.f_thickness):
            w.setEnabled(on)

    def _on_line_selected(self):
        spec = self._selected_line()
        self._loading = True
        try:
            if spec is None:
                self._set_form_enabled(False)
                self.form_pages.setCurrentWidget(self.form_placeholder)
                self.form_box.setTitle("Selected line")
                return
            self._set_form_enabled(True)
            self.form_box.setTitle(f"{TYPE_LABELS[spec.part_type]}: {spec.name}")
            self.f_name.setText(spec.name)
            self.f_name.setReadOnly(spec.from_parts)
            self.f_qty.setValue(spec.quantity)
            self.f_material.setText(spec.material)
            self.f_thickness.setValue(spec.thickness or 0)
            self.f_thickness.setEnabled(spec.part_type in ("sheet", "assembly"))
            if spec.part_type == "sheet":
                self.f_bends.setValue(-1 if spec.bends is None else spec.bends)
                self.f_tapped.setValue(spec.tapped_holes)
                self.f_tap_size.setCurrentText(spec.tap_size)
                self.f_paint.setChecked(spec.paint)
                try:
                    f = spec.to_features()
                    self.f_geom.setText(f"Flat {f.bbox[0]:.1f} x {f.bbox[1]:.1f} mm, net area "
                                        f"{f.net_area / 100:.1f} cm², cut length {f.cut_length:.0f} mm, "
                                        f"{f.pierces} pierce(s).")
                except ValueError as e:
                    self.f_geom.setText(str(e))
                self.form_pages.setCurrentWidget(self.page_sheet)
            elif spec.part_type == "tube":
                self.f_kind.setCurrentText(spec.profile_kind)
                for w, v in ((self.f_width, spec.width), (self.f_height, spec.height), (self.f_wall, spec.wall),
                             (self.f_length, spec.length), (self.f_bar, spec.bar_length),
                             (self.f_price_m, spec.price_per_m)):
                    w.setValue(v)
                self.f_end_a.setCurrentText(spec.end_a)
                self.f_end_b.setCurrentText(spec.end_b)
                self.f_tube_paint.setChecked(spec.paint)
                self.form_pages.setCurrentWidget(self.page_tube)
            elif spec.part_type == "purchased":
                self.f_unit_cost.setValue(spec.unit_cost)
                self.f_supplier.setText(spec.supplier)
                self.f_own_markup.setChecked(spec.markup_pct is not None)
                self.f_markup.setValue(spec.markup_pct if spec.markup_pct is not None else 0.0)
                self.f_markup.setEnabled(spec.markup_pct is not None)
                self.form_pages.setCurrentWidget(self.page_purchased)
            else:
                self.f_parts.setValue(spec.part_count)
                self.f_fasteners.setValue(spec.fastener_count)
                self.f_weld.setValue(spec.weld_length)
                self.form_pages.setCurrentWidget(self.page_assembly)
        finally:
            self._loading = False

    def _form_changed(self, *_):
        if self._loading:
            return
        spec = self._selected_line()
        if spec is None:
            return
        name = self.f_name.text().strip()
        if name and name != spec.name and not spec.from_parts:
            spec.name = self._unique_name(name)
        spec.quantity = self.f_qty.value()
        spec.material = self.f_material.text().strip()
        spec.thickness = self.f_thickness.value() or None
        if spec.part_type == "sheet":
            spec.bends = None if self.f_bends.value() < 0 else self.f_bends.value()
            spec.tapped_holes = self.f_tapped.value()
            spec.tap_size = self.f_tap_size.currentText()
            spec.paint = self.f_paint.isChecked()
        elif spec.part_type == "tube":
            spec.profile_kind = self.f_kind.currentText()
            spec.width, spec.height, spec.wall = self.f_width.value(), self.f_height.value(), self.f_wall.value()
            spec.length, spec.bar_length = self.f_length.value(), self.f_bar.value()
            spec.price_per_m = self.f_price_m.value()
            spec.end_a, spec.end_b = self.f_end_a.currentText(), self.f_end_b.currentText()
            spec.paint = self.f_tube_paint.isChecked()
        elif spec.part_type == "purchased":
            spec.unit_cost = self.f_unit_cost.value()
            spec.supplier = self.f_supplier.text().strip()
            self.f_markup.setEnabled(self.f_own_markup.isChecked())
            spec.markup_pct = self.f_markup.value() if self.f_own_markup.isChecked() else None
        else:
            spec.part_count = self.f_parts.value()
            spec.fastener_count = self.f_fasteners.value()
            spec.weld_length = self.f_weld.value()
        row = self.lines.index(spec)
        vals = (spec.name, TYPE_LABELS[spec.part_type], str(spec.quantity), self._material_text(spec),
                spec.summary())
        for c, v in enumerate(vals):
            self.lines_table.item(row, c).setText(v)
        self.form_box.setTitle(f"{TYPE_LABELS[spec.part_type]}: {spec.name}")

    # ------------------------------------------------------------ settings

    def read_setup(self) -> QuoteSetup:
        s = QuoteSetup(breaks=parse_breaks(self.s_breaks.text()),
                       material_mode=self.s_mode.currentData(), allocation=self.s_alloc.currentData(),
                       tube_machine=self.s_tube.currentData(),
                       paint_rate=self.s_paint.value(), transport=self.s_transport.value(),
                       fallback_price_per_kg=self.s_price_kg.value(),
                       kerf=self.panel.part_spacing.value())
        s.markups = {}
        for i, area in enumerate(COST_AREAS):
            text = (self.s_markups.item(i, 0).text() if self.s_markups.item(i, 0) else "").strip()
            try:
                s.markups[area] = float(text.rstrip("%")) if text else 0.0
            except ValueError:
                raise ValueError(f"markup for {area} isn't a number: {text!r}") from None
        return s

    def _refresh_rates_label(self):
        n_uncal = len(self.rates.uncalibrated_rows())
        n = len(self.rates.all_rows())
        tail = f" -- {n_uncal} of {n} rows uncalibrated" if n_uncal else f" -- {n} rows"
        self.rates_label.setText("Rates: " + self.rates_source + tail)

    def load_rates_file(self, path=None):
        if path is None:
            path, _ = QtWidgets.QFileDialog.getOpenFileName(
                self, "Load rates", "", "Rate tables (*.xlsx *.xlsm *.db *.sqlite);;All files (*)")
            if not path:
                return
        try:
            table = load_rates(path)
        except Exception as e:  # noqa: BLE001
            QtWidgets.QMessageBox.warning(self, "Quote", f"Couldn't load rates from {path}:\n{e}")
            return
        if not table.all_rows():
            QtWidgets.QMessageBox.warning(self, "Quote", f"{path} has no rate rows.")
            return
        self.rates = table
        self.rates_source = os.path.basename(path)
        self._refresh_rates_label()

    def export_rates_file(self, path=None):
        if path is None:
            path, _ = QtWidgets.QFileDialog.getSaveFileName(self, "Export rates", "rates.xlsx",
                                                            "Excel workbook (*.xlsx)")
            if not path:
                return
        if not path.lower().endswith((".xlsx", ".xlsm")):
            path += ".xlsx"
        try:
            save_rates(path, self.rates)
        except Exception as e:  # noqa: BLE001
            QtWidgets.QMessageBox.warning(self, "Quote", f"Couldn't write {path}:\n{e}")
            return
        self.status_label.setText(f"Rates written to {path} -- edit them there, then Load Rates.")

    # ------------------------------------------------------------- pricing

    @property
    def busy(self):
        return self._worker is not None

    def price(self):
        if self.busy:
            return
        if not self.lines:
            self.status_label.setText("Nothing to quote -- add parts on the Parts tab, or add a line.")
            return
        try:
            setup = self.read_setup()
        except ValueError as e:
            QtWidgets.QMessageBox.warning(self, "Quote", str(e))
            return
        lines = copy.deepcopy(self.lines)
        stock = copy.deepcopy(self.panel._inventory_template) if self.panel._inventory_template else None
        self._worker = _PriceWorker(lines, setup, self.rates, stock, self)
        self._worker.priced.connect(self._on_priced)
        self._worker.failed.connect(self._on_failed)
        self._worker.finished.connect(self._on_worker_done)
        self.status_label.setText(f"Pricing {len(lines)} line(s) at {', '.join(map(str, setup.breaks))} "
                                  f"set(s)...")
        self.busy_changed.emit(True)
        self._worker.start()

    def wait(self, ms=60000):
        """Block until a running price finishes (tests / scripted use)."""
        if self._worker is not None:
            self._worker.wait(ms)
            QtWidgets.QApplication.processEvents()

    def _on_worker_done(self):
        self._worker = None
        self.busy_changed.emit(False)

    def _on_failed(self, message):
        self.status_label.setText("Pricing failed -- see the message.")
        QtWidgets.QMessageBox.warning(self, "Quote", f"Pricing failed:\n{message}")

    def _on_priced(self, result, notes):
        self.result, self.result_notes = result, list(notes)
        # A re-price keeps overrides for lines/cells that still exist.
        for br in result.breaks:
            for line in br.lines:
                for key, cell in line.cells.items():
                    v = self.overrides.get((br.quantity, line.name, key))
                    if v is not None:
                        cell.override = v
            for key, cell in br.quote_cells.items():
                v = self.overrides.get((br.quantity, QUOTE, key))
                if v is not None:
                    cell.override = v
        self._fill_results()
        self.quote_ready.emit()

    # ------------------------------------------------------------- results

    def _money(self, v):
        import nesting_widgets
        sym = nesting_widgets.CURRENCY_SYMBOLS.get(self.panel.currency_code, "")
        return "—" if v is None else f"{sym}{v:,.2f}"

    def _fill_results(self):
        res = self.result
        if res is None:
            return
        keep = self._current_break_index()
        self.summary.setRowCount(len(res.breaks))
        for r, br in enumerate(res.breaks):
            margin = (br.total_price - br.total_cost) / br.total_price if br.total_price else 0.0
            vals = (str(br.quantity), self._money(br.total_cost), self._money(br.total_price),
                    self._money(br.set_price), f"{margin:.0%}",
                    "Complete" if br.complete else "Incomplete -- see notes")
            for c, v in enumerate(vals):
                item = QtWidgets.QTableWidgetItem(v)
                item.setTextAlignment(QtCore.Qt.AlignCenter)
                self.summary.setItem(r, c, item)
        self.summary.selectRow(min(keep, len(res.breaks) - 1) if keep is not None else 0)
        self.summary.resizeRowsToContents()
        h = self.summary.horizontalHeader().height() + 2 * self.summary.frameWidth() + sum(
            self.summary.rowHeight(r) for r in range(self.summary.rowCount()))
        self.summary.setFixedHeight(min(h + 2, 260))
        uncal = len(self.rates.uncalibrated_rows())
        self.status_label.setText(
            f"{len(res.breaks)} quantity break(s) priced." +
            (f"  {uncal} rate rows are uncalibrated starter values -- treat prices as estimates."
             if uncal else ""))
        self._fill_grid()

    def _current_break_index(self):
        rows = {i.row() for i in self.summary.selectedIndexes()}
        return rows.pop() if rows else None

    def current_break(self):
        if self.result is None or not self.result.breaks:
            return None
        i = self._current_break_index()
        return self.result.breaks[i if i is not None else 0]

    def grid_columns(self, br):
        keys = []
        for line in br.lines:
            for k in line.cells:
                if k not in keys:
                    keys.append(k)
        if MATERIAL in keys:
            keys.remove(MATERIAL)
            keys.insert(0, MATERIAL)
        return keys

    def _fill_grid(self):
        br = self.current_break()
        if br is None:
            return
        self._filling_grid = True
        try:
            keys = self.grid_columns(br)
            heads = ["Line", "Type", "Qty"] + [k.capitalize() for k in keys] + \
                    ["Unit cost", "Unit price", "Line total"]
            quote_rows = list(br.quote_cells.items())
            self.grid.clear()
            self.grid.setColumnCount(len(heads))
            self.grid.setHorizontalHeaderLabels(heads)
            self.grid.setRowCount(len(br.lines) + len(quote_rows) + 1)
            bold = QtGui.QFont(self.grid.font())
            bold.setBold(True)

            def ro(text, align=QtCore.Qt.AlignLeft | QtCore.Qt.AlignVCenter):
                item = QtWidgets.QTableWidgetItem(text)
                item.setFlags(item.flags() & ~QtCore.Qt.ItemIsEditable)
                item.setTextAlignment(align)
                return item

            right = QtCore.Qt.AlignRight | QtCore.Qt.AlignVCenter
            for r, line in enumerate(br.lines):
                self.grid.setItem(r, 0, ro(line.name))
                self.grid.setItem(r, 1, ro(TYPE_LABELS.get(line.part_type, line.part_type)))
                self.grid.setItem(r, 2, ro(str(line.quantity), right))
                for c, key in enumerate(keys, start=3):
                    cell = line.cells.get(key)
                    if cell is None:
                        self.grid.setItem(r, c, ro(""))
                        continue
                    item = QtWidgets.QTableWidgetItem("—" if cell.value is None else f"{cell.value:,.2f}")
                    item.setTextAlignment(right)
                    item.setData(QtCore.Qt.UserRole, (line.name, key))
                    tip = [f"{key}: {cell.cost_area}, markup {cell.markup:.0%}"]
                    if cell.overridden:
                        item.setFont(bold)
                        tip.append(f"Overridden -- calculated {self._money(cell.calculated)}. "
                                   f"Clear the cell to go back to it.")
                    if cell.estimate:
                        tip.append("Estimate (per-feature machining).")
                    if cell.value is None:
                        tip.append("Couldn't be priced -- see notes. Type a value to fill it in.")
                    item.setToolTip("\n".join(tip))
                    self.grid.setItem(r, c, item)
                base = 3 + len(keys)
                self.grid.setItem(r, base, ro(self._money(line.unit_cost), right))
                self.grid.setItem(r, base + 1, ro(self._money(line.unit_price), right))
                total = ro(self._money(line.total_price), right)
                if not line.complete:
                    total.setToolTip("Incomplete: " + "; ".join(line.notes))
                self.grid.setItem(r, base + 2, total)
            r0 = len(br.lines)
            for i, (name, cell) in enumerate(quote_rows):
                r = r0 + i
                self.grid.setItem(r, 0, ro(name.capitalize()))
                self.grid.setItem(r, 1, ro("Per quote"))
                item = QtWidgets.QTableWidgetItem(f"{cell.value:,.2f}" if cell.value is not None else "—")
                item.setTextAlignment(right)
                item.setData(QtCore.Qt.UserRole, (QUOTE, name))
                if cell.overridden:
                    item.setFont(bold)
                    item.setToolTip(f"Overridden -- calculated {self._money(cell.calculated)}.")
                self.grid.setItem(r, len(heads) - 3, item)
                self.grid.setItem(r, len(heads) - 1, ro(self._money(cell.price), right))
            tr = r0 + len(quote_rows)
            self.grid.setItem(tr, 0, ro("Total"))
            self.grid.item(tr, 0).setFont(bold)
            tot = ro(self._money(br.total_price), right)
            tot.setFont(bold)
            self.grid.setItem(tr, len(heads) - 1, tot)
            self.grid.resizeColumnsToContents()
            self.grid_label.setText(f"Breakdown at {br.quantity} set(s) -- per-unit costs; type in a "
                                    f"cell to override it for this quantity.")
            notes = list(self.result.notes) + self.result_notes + [
                f"{line.name}: {n}" for line in br.lines for n in line.notes] + list(br.notes)
            self.notes.setPlainText("\n".join(dict.fromkeys(notes)))
        finally:
            self._filling_grid = False

    def _on_grid_edit(self, item):
        if self._filling_grid:
            return
        ref = item.data(QtCore.Qt.UserRole)
        br = self.current_break()
        if not ref or br is None:
            return
        line_name, key = ref
        text = item.text().strip().replace(",", "")
        cell = (br.quote_cells.get(key) if line_name == QUOTE
                else br.line(line_name).cells.get(key))
        if cell is None:
            return
        okey = (br.quantity, line_name, key)
        if text in ("", "—"):
            cell.override = None
            self.overrides.pop(okey, None)
        else:
            try:
                value = float(text.lstrip("".join(nesting_currency_symbols())))
            except ValueError:
                QtWidgets.QMessageBox.warning(self, "Quote", f"{item.text()!r} isn't a number.")
                QtCore.QTimer.singleShot(0, self._fill_grid)
                return
            cell.override = value
            self.overrides[okey] = value
        QtCore.QTimer.singleShot(0, self._fill_results)

    def set_override(self, line_name, key, value, quantity=None):
        """Scripted equivalent of typing into a grid cell (None clears)."""
        br = self.current_break() if quantity is None else self.result.at(quantity)
        cell = br.quote_cells.get(key) if line_name == QUOTE else br.line(line_name).cells[key]
        cell.override = value
        if value is None:
            self.overrides.pop((br.quantity, line_name, key), None)
        else:
            self.overrides[(br.quantity, line_name, key)] = value
        self._fill_results()

    # -------------------------------------------------------------- export

    def export_xlsx(self, path=None):
        if self.result is None:
            QtWidgets.QMessageBox.information(self, "Quote", "Price the quote first.")
            return
        if path is None:
            path, _ = QtWidgets.QFileDialog.getSaveFileName(self, "Export quote", "quote.xlsx",
                                                            "Excel workbook (*.xlsx)")
            if not path:
                return
        if not path.lower().endswith(".xlsx"):
            path += ".xlsx"
        try:
            write_quote_xlsx(path, self.result, self.result_notes)
        except Exception as e:  # noqa: BLE001
            QtWidgets.QMessageBox.warning(self, "Quote", f"Couldn't write {path}:\n{e}")
            return
        self.status_label.setText(f"Quote written to {path}.")
        return path


def nesting_currency_symbols():
    """Every currency symbol the app shows, so a pasted "₹12.50" parses."""
    import nesting_widgets
    return tuple(nesting_widgets.CURRENCY_SYMBOLS.values())


def write_quote_xlsx(path, result, extra_notes=()):
    """Summary sheet (one row per quantity break) plus one breakdown sheet
    per break -- the same numbers as the Quote tab, overrides included."""
    try:
        import openpyxl
    except ImportError:
        raise RuntimeError("Exporting a quote needs the openpyxl package (pip install openpyxl).") from None
    from openpyxl.styles import Font
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Summary"
    ws.append(["Sets", "Total cost", "Total price", "Price per set", "Complete"])
    for br in result.breaks:
        ws.append([br.quantity, round(br.total_cost, 2), round(br.total_price, 2), round(br.set_price, 2),
                   "Yes" if br.complete else "No"])
    ws.append([])
    for n in list(result.notes) + list(extra_notes):
        ws.append([n])
    for c in ws[1]:
        c.font = Font(bold=True)
    for br in result.breaks:
        sh = wb.create_sheet(f"{br.quantity} off"[:31])
        keys = []
        for line in br.lines:
            for k in line.cells:
                if k not in keys:
                    keys.append(k)
        sh.append(["Line", "Type", "Qty"] + keys + ["Unit cost", "Unit price", "Line total", "Notes"])
        for c in sh[1]:
            c.font = Font(bold=True)
        for line in br.lines:
            row = [line.name, line.part_type, line.quantity]
            for k in keys:
                cell = line.cells.get(k)
                row.append(None if cell is None or cell.value is None else round(cell.value, 4))
            row += [round(line.unit_cost, 4), round(line.unit_price, 4), round(line.total_price, 2),
                    "; ".join(line.notes)]
            sh.append(row)
        for name, cell in br.quote_cells.items():
            sh.append([name, "per quote", 1] + [None] * len(keys) +
                      [cell.value, cell.price, cell.price, ""])
        sh.append(["Total"] + [None] * (len(keys) + 4) + [round(br.total_price, 2)])
    wb.save(path)
