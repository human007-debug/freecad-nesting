"""
quoting/costing_sheet_xlsx.py
-----------------------------
Writes a TenderCosting (quoting/tender_costing.py) as an Excel workbook in
the shop's own costing-sheet layout, with LIVE formulas -- so the
workbook still behaves like the hand-made one: change a weight, a rate or
the margin in Excel and everything below recalculates.

  COSTING SHEET   header block; A. MATERIALS (weight x % x rate, scrap
                  credit); B. PROCESS (qty x rate); C..H OTHERS; I ALL
                  TOTAL, J MARGIN / CONTINGENCIES, K BASIC COST, L GST,
                  M ALL INCLUSIVE RATE, N S&I, O LPR, P RECOMMENDED QUOTE
                  PRICE; margin needed to meet the LPR; tender value;
                  sign-off boxes; quoted rate and its ex-GST value.
  CONSOLIDATED BOM  every BOM item: drawing, description, spec, size,
                  qty per set, net and gross weights, qty per sheet/bar.
  COMPARISON      (when re-costing an imported workbook) the workbook's
                  own figures next to the re-costed ones.
"""

from typing import Optional

from quoting.tender_costing import FIELD_LABELS, TenderCosting


def write_costing_workbook(path: str, tc: TenderCosting, original: Optional[dict] = None):
    try:
        import openpyxl
    except ImportError:
        raise RuntimeError("Writing a costing sheet needs the openpyxl package (pip install openpyxl).") from None
    from openpyxl.styles import Alignment, Border, Font, PatternFill, Side

    bold = Font(bold=True)
    title = Font(bold=True, size=14)
    head_fill = PatternFill("solid", fgColor="D9E1F2")
    total_fill = PatternFill("solid", fgColor="FFF2CC")
    final_fill = PatternFill("solid", fgColor="E2EFDA")
    thin = Side(style="thin", color="999999")
    box = Border(left=thin, right=thin, top=thin, bottom=thin)
    money = '#,##0.00'
    wrap = Alignment(wrap_text=True, vertical="top")

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "COSTING SHEET"
    for col, width in zip("ABCDEFGHIJK", (7, 30, 16, 20, 8, 14, 8, 13, 15, 17, 34)):
        ws.column_dimensions[col].width = width

    def cell(ref, value, font=None, fill=None, fmt=None, align=None):
        c = ws[ref]
        c.value = value
        if font:
            c.font = font
        if fill:
            c.fill = fill
        if fmt:
            c.number_format = fmt
        if align:
            c.alignment = align
        return c

    def boxed(row, cols="ABCDEFGHIJK"):
        for col in cols:
            ws[f"{col}{row}"].border = box

    cell("A1", "COSTING SHEET", title)
    ws.merge_cells("A1:K1")
    h = tc.header
    r = 2
    for attr, label in FIELD_LABELS:
        cell(f"A{r}", label, bold)
        cell(f"F{r}", getattr(h, attr) or "", align=wrap)
        ws.merge_cells(f"F{r}:K{r}")
        boxed(r)
        r += 1
    cell(f"A{r}", "QUANTITY IN NOS", bold)
    cell(f"F{r}", h.quantity, fmt="0.00")
    cell(f"G{r}", h.unit)
    qty_row = r
    boxed(r)
    r += 1
    cell(f"A{r}", "PRODUCT WEIGHT", bold)
    weight_row = r
    boxed(r)
    r += 1

    # ---------------------------------------------------------- materials
    cell(f"A{r}", "MATERIALS", bold)
    r += 1
    for col, text in zip("ABCDEFGHIJK", ("S NO", "SHAPE", "SIZE", "MATL", "UOM", "WEIGHT", "%", "CHARGED WT",
                                        "ALL.INCL. RATE", "INPUT RM VALUE", "REMARKS")):
        cell(f"{col}{r}", text, bold, head_fill)
    boxed(r)
    r += 1
    first = r
    for i, line in enumerate(tc.material_lines(), start=1):
        cell(f"A{r}", i)
        cell(f"B{r}", line.shape)
        cell(f"C{r}", line.size)
        cell(f"D{r}", line.material)
        cell(f"E{r}", "KGS")
        cell(f"F{r}", line.weight, fmt="0.000")
        cell(f"G{r}", line.factor, fmt="0.00")
        cell(f"H{r}", f"=F{r}*G{r}", fmt="0.000")
        cell(f"I{r}", line.rate, fmt=money)
        cell(f"J{r}", f"=H{r}*I{r}", fmt=money)
        cell(f"K{r}", line.remarks)
        boxed(r)
        r += 1
    scrap_kg = tc.scrap_kg()
    scrap_value = tc.scrap_credit()
    scrap_rate = scrap_value / scrap_kg if scrap_kg else 0
    cell(f"F{r}", scrap_kg, fmt="0.00")
    cell(f"I{r}", scrap_rate, fmt=money)
    cell(f"J{r}", f"=-F{r}*I{r}", fmt=money)
    cell(f"K{r}", f"SCRAP {scrap_kg:.1f} KG (gross - net weight)")
    boxed(r)
    scrap_row = r
    r += 1
    mat_row = r
    cell(f"A{r}", "A", bold, total_fill)
    cell(f"B{r}", "TOTAL MATERIAL COST", bold, total_fill)
    cell(f"F{r}", f"=SUM(F{first}:F{scrap_row - 1})", bold, total_fill, "0.000")
    cell(f"H{r}", f"=SUM(H{first}:H{scrap_row - 1})", bold, total_fill, "0.000")
    cell(f"J{r}", f"=SUM(J{first}:J{scrap_row})", bold, total_fill, money)
    cell(f"K{r}", f"=J{r}/F{weight_row}", bold, total_fill, '#,##0.00" /kg"')
    boxed(r)
    cell(f"F{weight_row}", f"=F{mat_row}", fmt="0.000")
    cell(f"G{weight_row}", "KG")
    r += 2

    # ----------------------------------------------------------- process
    cell(f"A{r}", "PROCESS", bold)
    r += 1
    for col, text in zip("ABFHIJK", ("S NO", "PROCESS DESCRIPTION", "UOM", "QTY", "RATE", "VALUE", "REMARKS")):
        cell(f"{col}{r}", text, bold, head_fill)
    for col in "CDEG":
        ws[f"{col}{r}"].fill = head_fill
    boxed(r)
    r += 1
    first = r
    for i, p in enumerate(tc.process_lines(), start=1):
        cell(f"A{r}", i)
        cell(f"B{r}", p.description)
        ws.merge_cells(f"B{r}:E{r}")
        cell(f"F{r}", p.uom)
        cell(f"H{r}", p.qty, fmt="0.###")
        cell(f"I{r}", p.rate, fmt=money)
        cell(f"J{r}", f"=H{r}*I{r}", fmt=money)
        remark = p.remarks
        if p.derived_qty is not None and abs(p.derived_qty - p.qty) > 1e-6:
            remark = (remark + "; " if remark else "") + f"BOM gives {p.derived_qty:g}"
        cell(f"K{r}", remark)
        boxed(r)
        r += 1
    proc_row = r
    cell(f"A{r}", "B", bold, total_fill)
    cell(f"B{r}", "TOTAL PROCESS COST", bold, total_fill)
    cell(f"J{r}", f"=SUM(J{first}:J{r - 1})" if r > first else 0, bold, total_fill, money)
    cell(f"K{r}", f"=J{r}/F{weight_row}", bold, total_fill, '#,##0.00" /kg"')
    boxed(r)
    r += 2

    # ------------------------------------------------------------ others
    cell(f"A{r}", "OTHERS", bold)
    r += 1
    for col, text in zip("ABJ", ("S NO", "ITEM DESCRIPTION", "VALUE")):
        cell(f"{col}{r}", text, bold, head_fill)
    boxed(r)
    r += 1
    other_rows = []
    for o in tc.others:
        cell(f"A{r}", o.code)
        cell(f"B{r}", o.item)
        cell(f"J{r}", o.amount, fmt=money)
        cell(f"K{r}", o.remarks)
        boxed(r)
        other_rows.append(r)
        r += 1

    # ------------------------------------------------------------ totals
    t_row = r
    parts = [f"J{mat_row}", f"J{proc_row}"] + [f"J{x}" for x in other_rows]
    cell(f"A{r}", "I", bold, total_fill)
    cell(f"B{r}", "ALL TOTAL  A+B+C+D+E+F+G+H", bold, total_fill)
    cell(f"J{r}", "=" + "+".join(parts), bold, total_fill, money)
    boxed(r)
    r += 1
    m_row = r
    rate_row = r + 3
    cell(f"A{r}", "J", bold)
    cell(f"B{r}", tc.margin_label.upper(), bold)
    cell(f"F{r}", tc.margin_pct / 100.0, bold, PatternFill("solid", fgColor="FCE4D6"), "0.0%")
    cell(f"G{r}", "<- change the margin here")
    cell(f"I{r}", f"=J{r}/J{rate_row}", fmt="0.0%")
    cell(f"J{r}", f"=J{t_row}*F{r}", fmt=money)
    cell(f"K{r}", "margin % of total; % of all-inclusive rate in column I")
    boxed(r)
    r += 1
    basic_row = r
    cell(f"A{r}", "K", bold)
    cell(f"B{r}", "BASIC COST", bold)
    cell(f"J{r}", f"=J{t_row}+J{m_row}", bold, fmt=money)
    boxed(r)
    r += 1
    gst_row = r
    cell(f"A{r}", "L", bold)
    cell(f"B{r}", "GST", bold)
    cell(f"F{r}", tc.gst_pct / 100.0, fmt="0%")
    cell(f"J{r}", f"=J{basic_row}*F{r}", fmt=money)
    boxed(r)
    r += 1
    assert r == rate_row
    cell(f"A{r}", "M", bold, final_fill)
    cell(f"B{r}", "ALL INCLUSIVE RATE", bold, final_fill)
    cell(f"J{r}", f"=J{basic_row}+J{gst_row}", bold, final_fill, money)
    cell(f"K{r}", f"per {h.unit.lower()}", final_fill)
    boxed(r)
    r += 1
    cell(f"A{r}", "N", bold)
    cell(f"B{r}", "SUPPLY & INSTALLATION", bold)
    boxed(r)
    r += 1
    lpr_row = r
    cell(f"A{r}", "O", bold)
    cell(f"B{r}", "LPR", bold)
    cell(f"J{r}", tc.lpr, fmt=money)
    cell(f"K{r}", "last purchase rate")
    boxed(r)
    r += 1
    cell(f"A{r}", "P", bold, final_fill)
    cell(f"B{r}", "RECOMMENDED QUOTE PRICE", bold, final_fill)
    step = float(tc.card.setting("quote_rounding") or 1) or 1
    cell(f"J{r}", f"=ROUND(J{rate_row}/{step:g},0)*{step:g}", bold, final_fill, money)
    cell(f"K{r}", f"all-inclusive rate, rounded to {step:g}", final_fill)
    boxed(r)
    rec_row = r
    r += 2
    cell(f"B{r}", "MARGIN TO MEET LPR", bold)
    cell(f"J{r}", f'=IF(J{lpr_row}>0,(J{lpr_row}/(1+F{gst_row}))/J{t_row}-1,"")', bold, fmt="0.0%")
    cell(f"K{r}", "the margin at which M equals the LPR")
    r += 1
    cell(f"B{r}", f"TENDER VALUE ({h.quantity:g} {h.unit})", bold)
    cell(f"J{r}", f"=J{rec_row}*F{qty_row}", bold, fmt=money)
    cell(f"K{r}", "recommended price x quantity, incl. GST")
    r += 3
    cell(f"B{r}", "Functional Head", bold)
    cell(f"E{r}", "Costing Head", bold)
    cell(f"J{r}", "Director/President", bold)
    r += 3
    cell(f"H{r}", "QUOTED RATE:", bold)
    cell(f"J{r}", tc.quoted_rate if tc.quoted_rate is not None else f"=J{rec_row}", bold, fmt=money)
    r += 1
    cell(f"H{r}", "EX-GST:")
    cell(f"J{r}", f"=J{r - 1}/(1+F{gst_row})", fmt=money)
    ws.freeze_panes = "A2"
    ws.sheet_view.zoomScale = 90
    ws.page_setup.fitToWidth = 1

    # ------------------------------------------------------------- BOM
    bs = wb.create_sheet("CONSOLIDATED BOM")
    heads = ("S NO", "DRAWING", "DESCRIPTION", "TYPE", "GRADE", "THICK", "WIDTH", "LNTH", "QTY/SET",
             "WT/PC", "NET WT", "QTY/SHEET", "GROSS WT/PC", "WT./SET")
    bs.append(["CONSOLIDATED BILL OF MATERIAL -- " + (h.product or "")])
    bs["A1"].font = title
    bs.append(list(heads))
    for c in bs[2]:
        c.font, c.fill, c.border = bold, head_fill, box
    for i, w in enumerate(tc.weights(), start=1):
        it = w.item
        bs.append([i, it.drawing, it.description, it.shape, it.grade, it.thickness,
                   it.width if not isinstance(it.width, float) else round(it.width, 2), it.length, it.qty,
                   round(w.net_pc, 4), round(w.net, 3), w.qty_per_stock, round(w.gross_pc, 4), round(w.gross, 3)])
    n = bs.max_row
    bs.append(["", "", "TOTAL", "", "", "", "", "", "", "", f"=SUM(K3:K{n})", "", "", f"=SUM(N3:N{n})"])
    for c in bs[bs.max_row]:
        c.font = bold
    for col, width in zip("ABCDEFGHIJKLMN", (6, 26, 30, 12, 18, 8, 10, 9, 9, 10, 11, 10, 12, 11)):
        bs.column_dimensions[col].width = width
    bs.freeze_panes = "A3"

    # ------------------------------------------------------ comparison
    if original is not None:
        cs = wb.create_sheet("COMPARISON")
        cs.append(["", "YOUR SHEET", "RE-COSTED", "DIFFERENCE", "NOTE"])
        for c in cs[1]:
            c.font, c.fill = bold, head_fill
        t = tc.totals()
        orig = original.get("totals", {})
        for label, key, note in (("A. Material cost", "material", "rate card rates & usage %; BOM weights"),
                                 ("B. Process cost", "process", "sheet quantities, rate card rates"),
                                 ("C-H. Others", "others", ""),
                                 ("I. All total", "all_total", ""),
                                 ("M. All inclusive rate", "all_inclusive", f"{tc.margin_label.lower()} "
                                                                            f"{tc.margin_pct:.1f}% kept")):
            mine, theirs = t[key], orig.get(key)
            cs.append([label, theirs, round(mine, 2), None if theirs is None else round(mine - theirs, 2), note])
        cs.append([])
        cs.append(["Gross weight (kg)", original.get("gross_weight"), round(tc.gross_weight, 2)])
        cs.append(["Net weight (kg)", original.get("net_weight"), round(tc.net_weight, 2)])
        if tc.lpr:
            cs.append(["Margin to meet the LPR", None, round(tc.margin_for_rate(tc.lpr), 2), None,
                       f"LPR {tc.lpr:,.0f}"])
        for row in cs.iter_rows(min_row=2):
            for c in row[1:4]:
                c.number_format = money
        for col, width in zip("ABCDE", (26, 16, 16, 16, 50)):
            cs.column_dimensions[col].width = width
    wb.save(path)
