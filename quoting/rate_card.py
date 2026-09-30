"""
quoting/rate_card.py
--------------------
A shop's RATE CARD: the per-unit prices a costing sheet is built from,
the way a tender costing is actually done by hand -- material in currency
per kg by grade, shape and thickness; processes per metre, stroke, piece,
kg or sq ft; fixed "others" (packing, inspection, lab test ...); margin
and GST -- plus the sheet formula the BOM weights come from.

This is different from rates.py's machine tables (speed x hourly rate):
a rate card prices the work the way the quote is written, not the way the
machine runs. Both stay available.

THE CARD IS LOCAL. A shop's rates and margins are commercially sensitive,
and this repository may be public, so the real card is an .xlsx kept
outside the repo -- by default next to AlphaQuote's database
(`default_rate_card_path()`), or wherever ALPHAQUOTE_RATE_CARD points --
and .gitignore keeps rate-card files out of commits. The repo only ships
`example_rate_card()`: the same structure with placeholder values, for
tests and as a template.

Workbook layout (one sheet each; headings matched case/space-insensitively,
extra columns ignored so notes can be added freely):

  Settings    key | value | note
  Materials   material | grades | shape | thk_from | thk_to | rate_per_kg |
              usage_factor | scrap_per_kg | note
  Processes   process | uom | rate | basis | thk_from | thk_to | note
              (basis "fixed", or "per_mm" = rate x thickness in mm -- the
              laser-cutting rule "thickness x 3 per metre")
  Others      code | item | amount | note
  Fasteners   item | spec | price | note

SHEET FORMULA (`SheetFormula`, from the Settings sheet) -- how the BOM
turns a part into kilograms:
  weight per piece  = t x W x L x density
  sheet weight      = t x (sheet_w + allowance) x (sheet_l + allowance) x density
  strip weight      = t x (sheet_w + allowance) x (L + strip_allowance) x density
                      (a long part cut from its own full-width strip)
  qty per sheet     = floor(sheet_w / (W + gap_w)) x floor(sheet_l / (L + gap_l))
                      with gap_l = t when `length_gap_is_thickness`
  gross per piece   = sheet weight / qty per sheet
"""

import math
import os
from dataclasses import dataclass, field, fields
from typing import Dict, List, Optional

# Grade names as they appear on drawings and sheets -> one family key.
GRADE_ALIASES = {
    "ss409m": ("409m", "ss409m", "ss 409m", "x2crni12", "ss409", "409"),
    "ss304": ("304", "ss304", "ss 304", "aisi 304", "aisi304", "x5crni1810", "x5crni18-10",
              "ck 304", "ck304", "ss 304-2b", "ss 304 -2b", "ss 304 - 2b", "ss304-2b", "304-2b"),
    "ss410": ("410", "ss410", "ss 410", "x04cr12", "x04cr12 (410)"),
    "ss316": ("316", "ss316", "ss 316", "aisi 316", "x5crnimo17-12-2"),
    "ms": ("ms", "mild steel", "is2062", "is 2062", "e250", "s275"),
}

# Shapes as written on a costing sheet -> the rate card's shape classes.
SHAPE_CLASSES = {
    "sheet": ("sheet",),
    "plate": ("plate",),
    "tube": ("tube", "pipe", "sq tube", "sq. tube", "square tube", "rec tube", "rect tube",
             "rectangular tube", "shs", "rhs", "chs"),
    "bar": ("rod", "round", "round bar", "sq rod", "sq. rod", "square rod", "flat", "flat bar", "bar",
            "angle", "channel"),
}

SETTINGS_DEFAULTS = {
    "company": "",
    "currency": "INR",
    "density_g_cm3": 8.0,
    "sheet_width": 1250.0,
    "sheet_length": 2500.0,
    "sheet_allowance": 10.0,
    "strip_allowance": 20.0,
    "width_gap": 5.0,
    "length_gap_is_thickness": True,
    "tube_cut_allowance": 10.0,
    "bar_length": 6000.0,
    "tube_perimeter": "outer",
    "margin_pct": 12.0,
    "margin_label": "MARGIN",
    "gst_pct": 18.0,
    "quote_rounding": 100.0,
}


def _norm(text) -> str:
    return " ".join(str(text or "").strip().lower().replace("_", " ").split())


def grade_family(grade: Optional[str]) -> Optional[str]:
    """'X5CrNi1810' / 'AISI 304' / 'SS 304 -2B' -> 'ss304'. Unknown grades
    come back normalised, unchanged otherwise."""
    if not grade:
        return None
    g = _norm(grade)
    for family, names in GRADE_ALIASES.items():
        if g == family or g in names:
            return family
    # tolerate decorations: "C- K 201 X2CrNi12", "X2CrNi12 (SS409m)"
    compact = g.replace(" ", "")
    for family, names in GRADE_ALIASES.items():
        for n in sorted(names, key=len, reverse=True):
            if len(n) >= 4 and n.replace(" ", "") in compact:
                return family
    return g


def shape_class(shape: Optional[str]) -> str:
    s = _norm(shape)
    for cls, names in SHAPE_CLASSES.items():
        if s == cls or s in names:
            return cls
    for cls, names in SHAPE_CLASSES.items():
        if any(n in s for n in names):
            return cls
    return s or "sheet"


# ------------------------------------------------------------------ rows

@dataclass
class MaterialRate:
    material: str                     # family key or display name, e.g. "SS304"
    grades: str = ""                  # comma-separated names that mean this material
    shape: str = "sheet"              # sheet / plate / tube / bar
    thk_from: float = 0.0             # mm, inclusive
    thk_to: float = 1000.0            # mm, inclusive
    rate_per_kg: float = 0.0
    usage_factor: float = 1.0         # the costing sheet's "%" column
    scrap_per_kg: float = 0.0
    note: str = ""

    def matches(self, family: str, shape: str, thickness: Optional[float]) -> bool:
        names = {grade_family(self.material)} | {grade_family(g) for g in self.grades.split(",") if g.strip()}
        if family not in names or shape_class(self.shape) != shape:
            return False
        if thickness is None:
            return True
        return self.thk_from - 1e-9 <= thickness <= self.thk_to + 1e-9


@dataclass
class ProcessRate:
    process: str
    uom: str
    rate: float
    basis: str = "fixed"              # fixed | per_mm
    thk_from: float = 0.0
    thk_to: float = 1000.0
    note: str = ""

    def rate_for(self, thickness: Optional[float] = None) -> float:
        if self.basis == "per_mm":
            if thickness is None:
                raise ValueError(f"{self.process} is priced per mm of thickness -- needs a thickness")
            return self.rate * thickness
        return self.rate


@dataclass
class OtherCharge:
    code: str
    item: str
    amount: float = 0.0
    note: str = ""


@dataclass
class FastenerPrice:
    item: str
    spec: str = ""
    price: float = 0.0
    note: str = ""


class RateCardError(LookupError):
    pass


@dataclass
class SheetFormula:
    density_g_cm3: float = 8.0
    sheet_width: float = 1250.0
    sheet_length: float = 2500.0
    sheet_allowance: float = 10.0
    strip_allowance: float = 20.0
    width_gap: float = 5.0
    length_gap_is_thickness: bool = True
    tube_cut_allowance: float = 10.0
    tube_perimeter: str = "outer"

    @property
    def kg_per_mm3(self) -> float:
        return self.density_g_cm3 / 1e6

    def piece_weight(self, t, w, l) -> float:
        return t * w * l * self.kg_per_mm3

    def sheet_weight(self, t) -> float:
        return t * (self.sheet_width + self.sheet_allowance) * (self.sheet_length + self.sheet_allowance) \
            * self.kg_per_mm3

    def strip_weight(self, t, l) -> float:
        return t * (self.sheet_width + self.sheet_allowance) * (l + self.strip_allowance) * self.kg_per_mm3

    def qty_per_sheet(self, t, w, l) -> int:
        gap_l = t if self.length_gap_is_thickness else 0.0
        return int(math.floor(self.sheet_width / (w + self.width_gap))) * \
            int(math.floor(self.sheet_length / (l + gap_l)))

    def gross_weight(self, t, w, l, qty_per_sheet: Optional[int] = None) -> float:
        """Material charged for one piece: its share of a sheet, or of a
        full-width strip when not even one fits along the sheet's length."""
        if l + (t if self.length_gap_is_thickness else 0) > self.sheet_length:
            per_strip = max(1, int(math.floor(self.sheet_width / (w + self.width_gap))))
            return self.strip_weight(t, l) / per_strip
        n = qty_per_sheet if qty_per_sheet else self.qty_per_sheet(t, w, l)
        if n < 1:
            raise ValueError(f"a {w:g} x {l:g} mm piece doesn't fit a "
                             f"{self.sheet_width:g} x {self.sheet_length:g} sheet")
        return self.sheet_weight(t) / n

    def square_tube_weight(self, t, side, l) -> float:
        """Perimeter x wall x length -- outer perimeter (4 x side), or the
        mean one (4 x (side - wall)) with `tube_perimeter = "mean"`."""
        return self.rect_tube_weight(t, side, side, l)

    def rect_tube_weight(self, t, a, b, l) -> float:
        per = 2 * (a + b) - (4 * t if self.tube_perimeter == "mean" else 0)
        return t * per * l * self.kg_per_mm3

    def round_tube_weight(self, t, od, l) -> float:
        return math.pi * od * t * l * self.kg_per_mm3

    def bar_weight(self, dia, l) -> float:
        return math.pi / 4 * dia * dia * l * self.kg_per_mm3


@dataclass
class RateCard:
    settings: Dict[str, object] = field(default_factory=lambda: dict(SETTINGS_DEFAULTS))
    materials: List[MaterialRate] = field(default_factory=list)
    processes: List[ProcessRate] = field(default_factory=list)
    others: List[OtherCharge] = field(default_factory=list)
    fasteners: List[FastenerPrice] = field(default_factory=list)
    source: str = ""

    def setting(self, key):
        return self.settings.get(key, SETTINGS_DEFAULTS.get(key))

    @property
    def formula(self) -> SheetFormula:
        kw = {}
        for f in fields(SheetFormula):
            v = self.setting(f.name)
            if isinstance(f.default, bool):
                kw[f.name] = _as_bool(v)
            elif isinstance(f.default, str):
                kw[f.name] = _norm(v) or f.default
            else:
                kw[f.name] = float(v)
        return SheetFormula(**kw)

    def material(self, grade: str, shape: str = "sheet", thickness: Optional[float] = None) -> MaterialRate:
        fam, cls = grade_family(grade), shape_class(shape)
        hits = [m for m in self.materials if m.matches(fam, cls, thickness)]
        if not hits and cls == "plate":          # a card without plate rows prices plate as sheet
            hits = [m for m in self.materials if m.matches(fam, "sheet", thickness)]
        if not hits:
            raise RateCardError(f"no rate for {grade} {shape}" +
                                (f" {thickness:g} mm" if thickness is not None else "") + " on the rate card")
        return min(hits, key=lambda m: m.thk_to - m.thk_from)   # the narrowest band wins

    def process(self, name: str, thickness: Optional[float] = None) -> ProcessRate:
        key = _norm(name)
        hits = [p for p in self.processes if _norm(p.process) == key]
        if not hits:
            raise RateCardError(f"no process {name!r} on the rate card")
        if thickness is not None:
            banded = [p for p in hits if p.thk_from - 1e-9 <= thickness <= p.thk_to + 1e-9]
            if banded:
                hits = banded
        return min(hits, key=lambda p: p.thk_to - p.thk_from)

    def process_rate(self, name: str, thickness: Optional[float] = None) -> float:
        return self.process(name, thickness).rate_for(thickness)

    def other(self, code_or_item: str) -> OtherCharge:
        key = _norm(code_or_item)
        for o in self.others:
            if _norm(o.code) == key or _norm(o.item) == key:
                return o
        raise RateCardError(f"no charge {code_or_item!r} on the rate card")

    def summary(self) -> str:
        who = self.setting("company") or "Rate card"
        return (f"{who}: {len(self.materials)} material rates, {len(self.processes)} processes, "
                f"{len(self.others)} other charges, {len(self.fasteners)} fasteners")


def _as_bool(v) -> bool:
    if isinstance(v, bool):
        return v
    return _norm(v) in ("1", "y", "yes", "true", "t", "x")


# ----------------------------------------------------------------- storage

_SHEETS = (("Materials", MaterialRate, "materials"), ("Processes", ProcessRate, "processes"),
           ("Others", OtherCharge, "others"), ("Fasteners", FastenerPrice, "fasteners"))


def _require_openpyxl():
    try:
        import openpyxl
    except ImportError:
        raise RuntimeError("Reading/writing a rate card needs the openpyxl package "
                           "(pip install openpyxl).") from None
    return openpyxl


def _coerce(f, value):
    if value is None or (isinstance(value, str) and not value.strip()):
        return f.default
    if f.type in (float, "float"):
        return float(str(value).replace(",", ""))
    if f.type in (str, "str"):
        return str(value).strip()
    return value


def save_rate_card(path: str, card: RateCard):
    openpyxl = _require_openpyxl()
    from openpyxl.styles import Font, PatternFill
    head_font, head_fill = Font(bold=True, color="FFFFFF"), PatternFill("solid", fgColor="1F4E78")
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Settings"
    ws.append(["key", "value", "note"])
    notes = {
        "density_g_cm3": "used for every weight on the BOM",
        "sheet_width": "stock sheet, mm", "sheet_length": "stock sheet, mm",
        "sheet_allowance": "added to both sheet sides when weighing a sheet (1250 -> 1260)",
        "strip_allowance": "added to a long part's length when it takes its own strip",
        "width_gap": "gap between parts across the sheet when counting qty/sheet",
        "length_gap_is_thickness": "gap along the sheet = material thickness",
        "tube_cut_allowance": "added to each tube length",
        "bar_length": "stock length rods and flats are bought in, mm",
        "tube_perimeter": "outer (4 x side x wall) or mean (4 x (side - wall) x wall) for square/rect tube",
        "margin_pct": "default margin on ALL TOTAL -- adjusted per quote",
        "margin_label": "MARGIN or CONTINGENCIES",
        "gst_pct": "GST on BASIC COST",
        "quote_rounding": "quoted rate rounded to this",
    }
    for key in list(SETTINGS_DEFAULTS) + [k for k in card.settings if k not in SETTINGS_DEFAULTS]:
        v = card.setting(key)
        ws.append([key, ("Yes" if v else "No") if isinstance(v, bool) else v, notes.get(key, "")])
    for title, cls, attr in _SHEETS:
        sh = wb.create_sheet(title)
        cols = [f.name for f in fields(cls)]
        sh.append(cols)
        for row in getattr(card, attr):
            sh.append([getattr(row, c) for c in cols])
    for sh in wb.worksheets:
        for c in sh[1]:
            c.font, c.fill = head_font, head_fill
        sh.freeze_panes = "A2"
        for col in sh.columns:
            width = max(len(str(c.value)) if c.value is not None else 0 for c in col)
            sh.column_dimensions[col[0].column_letter].width = min(max(10, width + 2), 60)
    wb.save(path)


def load_rate_card(path: str) -> RateCard:
    openpyxl = _require_openpyxl()
    wb = openpyxl.load_workbook(path, data_only=True, read_only=True)
    card = RateCard(settings=dict(SETTINGS_DEFAULTS), source=path)
    try:
        names = {_norm(n): n for n in wb.sheetnames}
        if "settings" in names:
            for row in list(wb[names["settings"]].iter_rows(values_only=True))[1:]:
                if not row or row[0] is None:
                    continue
                key = _norm(row[0]).replace(" ", "_")
                value = row[1] if len(row) > 1 else None
                default = SETTINGS_DEFAULTS.get(key)
                if isinstance(default, bool):
                    value = _as_bool(value)
                elif isinstance(default, str) and value is not None:
                    value = str(value).strip()
                elif isinstance(default, float) and value is not None:
                    value = float(value)
                card.settings[key] = value if value is not None else default
        for title, cls, attr in _SHEETS:
            if _norm(title) not in names:
                continue
            rows = list(wb[names[_norm(title)]].iter_rows(values_only=True))
            if not rows:
                continue
            heads = [_norm(h).replace(" ", "_") for h in rows[0]]
            by_name = {f.name: f for f in fields(cls)}
            first = fields(cls)[0].name
            for i, raw in enumerate(rows[1:], start=2):
                data = {h: v for h, v in zip(heads, raw) if h in by_name}
                if data.get(first) in (None, ""):
                    continue
                try:
                    getattr(card, attr).append(cls(**{n: _coerce(f, data.get(n)) for n, f in by_name.items()}))
                except (TypeError, ValueError) as e:
                    raise ValueError(f"{path}: sheet {title!r} row {i}: {e}") from None
    finally:
        wb.close()
    return card


def default_rate_card_path() -> str:
    """ALPHAQUOTE_RATE_CARD if set, else rate_card.xlsx next to AlphaQuote's
    database (quoting/db.py default_db_path())."""
    env = os.environ.get("ALPHAQUOTE_RATE_CARD")
    if env:
        return env
    from quoting.db import default_db_path
    return os.path.join(os.path.dirname(default_db_path()), "rate_card.xlsx")


def example_rate_card() -> RateCard:
    """The structure of a rate card with PLACEHOLDER values -- what ships
    in the repo. Replace every number with the shop's own."""
    note = "example -- replace"
    return RateCard(
        settings={**SETTINGS_DEFAULTS, "company": "Example shop"},
        materials=[
            MaterialRate("SS304", "AISI 304, X5CrNi1810", "sheet", 0, 6, 180.0, 0.95, 50.0, note),
            MaterialRate("SS304", "AISI 304, X5CrNi1810", "plate", 6, 50, 180.0, 1.0, 50.0, note),
            MaterialRate("SS304", "AISI 304, X5CrNi1810", "tube", 0, 20, 200.0, 0.9, 50.0, note),
            MaterialRate("SS304", "AISI 304, X5CrNi1810", "bar", 0, 200, 190.0, 1.0, 50.0, note),
            MaterialRate("MS", "IS2062", "sheet", 0, 50, 70.0, 0.95, 20.0, note),
        ],
        processes=[
            ProcessRate("LASER CUTTING", "METRE", 2.0, "per_mm", 0, 10, note),
            ProcessRate("LASER CUTTING", "METRE", 4.0, "per_mm", 10, 50, note),
            ProcessRate("SHEARING", "KGS", 2.0, note=note),
            ProcessRate("BENDING", "STROKE", 5.0, note=note),
            ProcessRate("DRILLING", "NOS", 8.0, note=note),
            ProcessRate("WELDING", "METRE", 40.0, note=note),
            ProcessRate("PAINTING", "SQ. FT", 6.0, note=note),
        ],
        others=[
            OtherCharge("C", "BOUGHT OUT PARTS", 0.0, note),
            OtherCharge("D", "SUPPLY & INSTALLATION", 0.0, note),
            OtherCharge("E", "PACKING & TRANSPORT CHARGES", 1000.0, note),
            OtherCharge("F", "INSPECTION FEES", 500.0, note),
            OtherCharge("G", "OTHERS", 200.0, note),
            OtherCharge("H", "LAB TEST REPORT", 200.0, note),
        ],
        fasteners=[FastenerPrice("HEX BOLT (M8x60)", "IS 1364", 20.0, note)],
        source="example",
    )
