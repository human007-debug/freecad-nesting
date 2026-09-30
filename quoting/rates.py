"""
quoting/rates.py
----------------
Phase 2: machine and operation rate tables.

Each kind of rate is a small dataclass (one row = one material/thickness,
weld type/thickness, saw size band, ...). A `RateTable` holds lists of
rows of every kind and answers lookups:

  - `SIZE_MODE = "interpolate"` (cutting speed vs thickness, etc.): an
    exact size match, or linear interpolation between the two nearest
    sizes. A size outside the table raises `RateLookupError` rather than
    extrapolating -- a quote silently priced off the end of a speed chart
    is exactly the kind of error nobody notices until the job loses money.
  - `SIZE_MODE = "band"` (saw time by section size, drill time by
    diameter): the row's size is an upper bound; the smallest row that
    covers the value wins.

Rows are matched on a category first (material, weld type, feature kind),
normalised through features.normalize_material; a row whose category is
"*" is the fallback for any category without its own rows.

Storage mirrors inventory.py's stock list: SQLite for the app
(`save_rates_sqlite`/`load_rates_sqlite`, one table per rate kind) and
XLSX for people (`save_rates_xlsx`/`load_rates_xlsx`, one worksheet per
rate kind, field names as the heading row). openpyxl is imported lazily,
same as inventory.py.

`starter_rates()` ships round-number rates for mild steel, stainless and
aluminium fibre-laser cutting plus press-brake, weld, saw, tube-laser,
machining and assembly rows. Every one is `calibrated=False` with a
"calibrate me" note -- real speeds depend on the machine, gas and nozzle,
and these are only a plausible starting point (see the plan's risks: first
estimates may be 10-30% off until calibrated). `uncalibrated_rows()` lists
what's still on the starter values.
"""

import dataclasses
import sqlite3
from dataclasses import dataclass, fields
from typing import Dict, List, Optional, Type

from quoting.features import normalize_material

CALIBRATE_ME = "calibrate me"


class RateLookupError(LookupError):
    pass


@dataclass
class CuttingRate:
    """Flat-sheet laser cutting, per material and thickness."""
    material: str
    thickness: float          # mm
    speed_mm_min: float       # cutting speed
    pierce_s: float           # seconds per pierce
    rate_per_hr: float        # machine + operator, currency/hr
    setup_min: float = 10.0   # per batch: program load, sheet handling
    calibrated: bool = False
    note: str = ""
    CATEGORY = "material"
    SIZE = "thickness"
    SIZE_MODE = "interpolate"


@dataclass
class BendRate:
    """Press brake, per material ("*" = any) and thickness."""
    material: str
    thickness: float
    hit_s: float              # seconds per bend (cycle incl. handling)
    setup_min: float          # per batch: first tool set, first-off check
    tool_change_min: float    # each extra tool set a part needs
    rate_per_hr: float
    calibrated: bool = False
    note: str = ""
    CATEGORY = "material"
    SIZE = "thickness"
    SIZE_MODE = "interpolate"


@dataclass
class WeldRate:
    """Manual welding, per weld type and plate thickness."""
    weld_type: str            # fillet / butt / "*"
    thickness: float
    min_per_m: float          # arc + dressing minutes per metre of weld
    filler_per_m: float       # consumables, currency per metre
    tack_s: float             # seconds per tack (manual tack-count fallback)
    setup_min: float          # fixturing per assembly batch
    rate_per_hr: float
    calibrated: bool = False
    note: str = ""
    CATEGORY = "weld_type"
    SIZE = "thickness"
    SIZE_MODE = "interpolate"


@dataclass
class SawRate:
    """Cold/band saw, per section-size band (largest outer dimension)."""
    max_section: float        # mm -- this row covers sections up to this size
    cut_s: float              # seconds per square end
    setup_min: float
    rate_per_hr: float
    calibrated: bool = False
    note: str = ""
    CATEGORY = None
    SIZE = "max_section"
    SIZE_MODE = "band"


@dataclass
class TubeLaserRate:
    """Tube laser, per material and wall thickness."""
    material: str
    wall: float
    speed_mm_min: float
    pierce_s: float
    load_s: float             # per part: load, chuck, unload
    setup_min: float
    rate_per_hr: float
    calibrated: bool = False
    note: str = ""
    CATEGORY = "material"
    SIZE = "wall"
    SIZE_MODE = "interpolate"


@dataclass
class MachiningRate:
    """Per-feature machining, per feature kind and diameter band.
    `feature` is one of features.HOLE_KINDS or "pocket"; a pocket row uses
    `removal_cm3_min` instead of `time_s`."""
    feature: str
    max_diameter: float       # mm -- this row covers features up to this size
    time_s: float             # per feature
    removal_cm3_min: float    # pockets/slots only
    setup_min: float          # per machining setup
    rate_per_hr: float
    calibrated: bool = False
    note: str = ""
    CATEGORY = "feature"
    SIZE = "max_diameter"
    SIZE_MODE = "band"


@dataclass
class AssemblyRate:
    """Assembly labour (a single row, or one per `kind` of assembly)."""
    kind: str
    base_min: float           # per assembly
    per_part_min: float
    per_fastener_min: float
    rate_per_hr: float
    calibrated: bool = False
    note: str = ""
    CATEGORY = "kind"
    SIZE = None
    SIZE_MODE = None


RATE_KINDS: Dict[str, Type] = {
    cls.__name__: cls
    for cls in (CuttingRate, BendRate, WeldRate, SawRate, TubeLaserRate, MachiningRate, AssemblyRate)
}


def _norm_category(cls, value) -> Optional[str]:
    if value is None:
        return None
    if cls.CATEGORY == "material":
        return normalize_material(value)
    return str(value).strip().lower()


class RateTable:
    def __init__(self, rows: Optional[list] = None):
        self.rows: Dict[str, list] = {name: [] for name in RATE_KINDS}
        for row in rows or []:
            self.add(row)

    def add(self, row):
        name = type(row).__name__
        if name not in RATE_KINDS:
            raise TypeError(f"not a rate row: {row!r}")
        self.rows[name].append(row)
        return row

    def all_rows(self) -> list:
        return [r for rows in self.rows.values() for r in rows]

    def of_kind(self, cls) -> list:
        return list(self.rows[cls.__name__])

    def uncalibrated_rows(self) -> list:
        return [r for r in self.all_rows() if not r.calibrated]

    # ------------------------------------------------------------- lookup

    def _candidates(self, cls, category):
        rows = self.rows[cls.__name__]
        if cls.CATEGORY is None:
            return rows
        want = _norm_category(cls, category)
        exact = [r for r in rows if _norm_category(cls, getattr(r, cls.CATEGORY)) == want]
        if exact:
            return exact
        return [r for r in rows if getattr(r, cls.CATEGORY) == "*"]

    def lookup(self, cls, category=None, size: Optional[float] = None):
        """The row (or an interpolated copy of one) for `category` at
        `size`. See the module docstring for the matching rules."""
        rows = self._candidates(cls, category)
        what = f"{cls.__name__}" + (f" for {category!r}" if cls.CATEGORY else "")
        if not rows:
            raise RateLookupError(f"no {what} in the rate table")
        if cls.SIZE_MODE is None:
            return rows[0]
        if size is None:
            raise RateLookupError(f"{what} needs a {cls.SIZE}")
        key = cls.SIZE
        rows = sorted(rows, key=lambda r: getattr(r, key))

        if cls.SIZE_MODE == "band":
            for r in rows:
                if size <= getattr(r, key) + 1e-9:
                    return r
            raise RateLookupError(
                f"{what}: {cls.SIZE} {size:g} is above the largest band ({getattr(rows[-1], key):g})")

        for r in rows:
            if abs(getattr(r, key) - size) <= 1e-6:
                return r
        lo = [r for r in rows if getattr(r, key) < size]
        hi = [r for r in rows if getattr(r, key) > size]
        if not lo or not hi:
            span = f"{getattr(rows[0], key):g}-{getattr(rows[-1], key):g}"
            raise RateLookupError(f"{what}: {cls.SIZE} {size:g} is outside the table ({span}); "
                                  f"add a row rather than extrapolating")
        return _interpolate(lo[-1], hi[0], key, size)


def _interpolate(a, b, key, size):
    t = (size - getattr(a, key)) / (getattr(b, key) - getattr(a, key))
    values = {}
    for f in fields(a):
        va, vb = getattr(a, f.name), getattr(b, f.name)
        if f.name == key:
            values[f.name] = size
        elif isinstance(va, bool):
            values[f.name] = va and vb
        elif isinstance(va, (int, float)) and isinstance(vb, (int, float)):
            values[f.name] = va + (vb - va) * t
        elif f.name == "note":
            values[f.name] = va if va == vb else f"interpolated: {va} / {vb}".strip()
        else:
            values[f.name] = va
    return type(a)(**values)


# ------------------------------------------------------------------ storage

def _columns(cls):
    return [f for f in fields(cls)]


def _sql_type(f):
    if f.type in (float, "float"):
        return "REAL"
    if f.type in (bool, int, "bool", "int"):
        return "INTEGER"
    return "TEXT"


def _coerce(f, value):
    if value is None or value == "":
        if f.default is not dataclasses.MISSING:
            return f.default
        raise ValueError(f"{f.name} is required")
    t = f.type
    if t in (float, "float"):
        return float(value)
    if t in (int, "int"):
        return int(float(value))
    if t in (bool, "bool"):
        if isinstance(value, str):
            return value.strip().lower() in ("1", "y", "yes", "true", "t", "x")
        return bool(value)
    return str(value)


def save_rates_sqlite(path: str, table: RateTable):
    """Replace every rate table in the database at `path` with `table`."""
    con = sqlite3.connect(path)
    try:
        with con:
            for name, cls in RATE_KINDS.items():
                cols = _columns(cls)
                con.execute(f'DROP TABLE IF EXISTS "{name}"')
                con.execute(f'CREATE TABLE "{name}" (' +
                            ", ".join(f'"{f.name}" {_sql_type(f)}' for f in cols) + ")")
                marks = ", ".join("?" for _ in cols)
                con.executemany(
                    f'INSERT INTO "{name}" VALUES ({marks})',
                    [tuple(getattr(r, f.name) for f in cols) for r in table.rows[name]],
                )
    finally:
        con.close()


def load_rates_sqlite(path: str) -> RateTable:
    table = RateTable()
    con = sqlite3.connect(path)
    try:
        existing = {r[0] for r in con.execute("SELECT name FROM sqlite_master WHERE type='table'")}
        for name, cls in RATE_KINDS.items():
            if name not in existing:
                continue
            cur = con.execute(f'SELECT * FROM "{name}"')
            heads = [d[0] for d in cur.description]
            by_name = {f.name: f for f in _columns(cls)}
            for row in cur:
                data = dict(zip(heads, row))
                table.add(cls(**{n: _coerce(f, data.get(n)) for n, f in by_name.items()}))
    finally:
        con.close()
    return table


def _require_openpyxl():
    try:
        import openpyxl
    except ImportError:
        raise RuntimeError("Reading/writing an Excel rate file needs the openpyxl package "
                           "(pip install openpyxl).") from None
    return openpyxl


def save_rates_xlsx(path: str, table: RateTable):
    openpyxl = _require_openpyxl()
    wb = openpyxl.Workbook()
    wb.remove(wb.active)
    for name, cls in RATE_KINDS.items():
        ws = wb.create_sheet(name)
        cols = _columns(cls)
        ws.append([f.name for f in cols])
        for r in table.rows[name]:
            ws.append([getattr(r, f.name) for f in cols])
        ws.freeze_panes = "A2"
    wb.save(path)


def load_rates_xlsx(path: str) -> RateTable:
    """Worksheets not named after a rate kind are ignored; columns are
    matched by heading (case/space-insensitive), so a person can reorder
    them or add their own notes columns."""
    openpyxl = _require_openpyxl()
    wb = openpyxl.load_workbook(path, data_only=True, read_only=True)
    table = RateTable()
    try:
        for name, cls in RATE_KINDS.items():
            if name not in wb.sheetnames:
                continue
            rows = list(wb[name].iter_rows(values_only=True))
            if not rows:
                continue
            heads = [str(h or "").strip().lower().replace(" ", "_") for h in rows[0]]
            by_name = {f.name: f for f in _columns(cls)}
            for i, raw in enumerate(rows[1:], start=2):
                if all(v is None or str(v).strip() == "" for v in raw):
                    continue
                data = {h: v for h, v in zip(heads, raw) if h}
                try:
                    table.add(cls(**{n: _coerce(f, data.get(n)) for n, f in by_name.items()}))
                except ValueError as e:
                    raise ValueError(f"{path}: sheet {name!r} row {i}: {e}") from None
    finally:
        wb.close()
    return table


def save_rates(path: str, table: RateTable):
    if str(path).lower().endswith((".xlsx", ".xlsm")):
        save_rates_xlsx(path, table)
    else:
        save_rates_sqlite(path, table)


def load_rates(path: str) -> RateTable:
    if str(path).lower().endswith((".xlsx", ".xlsm")):
        return load_rates_xlsx(path)
    return load_rates_sqlite(path)


# ------------------------------------------------------------ starter rates

# (thickness mm, speed mm/min, pierce s) -- round numbers in the range a
# ~3-4 kW fibre laser manages (N2 on stainless/aluminium, O2 on thicker
# mild steel). NOT measured on any machine.
_LASER_STARTER = {
    "mild steel": [(1, 20000, 0.2), (2, 9000, 0.3), (3, 5000, 0.5), (5, 3000, 1.0),
                   (8, 1600, 2.0), (10, 1200, 3.0), (12, 900, 4.0)],
    "stainless": [(1, 22000, 0.2), (2, 9000, 0.3), (3, 4500, 0.5), (5, 1600, 1.2),
                  (8, 600, 2.5), (10, 400, 4.0)],
    "aluminium": [(1, 22000, 0.2), (2, 10000, 0.3), (3, 5000, 0.5), (5, 1500, 1.2),
                  (8, 500, 3.0)],
}
_LASER_RATE_PER_HR = 120.0


def starter_rates() -> RateTable:
    t = RateTable()
    for material, rows in _LASER_STARTER.items():
        for thk, speed, pierce in rows:
            t.add(CuttingRate(material, float(thk), float(speed), float(pierce),
                              _LASER_RATE_PER_HR, 10.0, False, CALIBRATE_ME))
    for thk, hit, setup in [(0.5, 10, 15), (3, 12, 20), (6, 18, 25), (12, 30, 35)]:
        t.add(BendRate("*", float(thk), float(hit), float(setup), 8.0, 75.0, False, CALIBRATE_ME))
    for wt, rows in {"fillet": [(1, 8, 1.0), (3, 10, 1.5), (6, 16, 3.0), (12, 40, 8.0)],
                     "butt": [(1, 10, 1.2), (3, 14, 2.0), (6, 25, 4.5), (12, 60, 12.0)]}.items():
        for thk, min_per_m, filler in rows:
            t.add(WeldRate(wt, float(thk), float(min_per_m), float(filler), 20.0, 20.0, 70.0,
                           False, CALIBRATE_ME))
    for size, cut_s in [(50, 30), (100, 60), (200, 150), (400, 360)]:
        t.add(SawRate(float(size), float(cut_s), 10.0, 60.0, False, CALIBRATE_ME))
    for material in ("mild steel", "stainless", "aluminium"):
        for wall, speed in [(1, 12000), (3, 4000), (6, 1500)]:
            t.add(TubeLaserRate(material, float(wall), float(speed), 0.5, 20.0, 15.0, 150.0,
                                False, CALIBRATE_ME))
    for feature, rows in {"plain": [(6, 15), (12, 25), (25, 45), (50, 120)],
                          "tapped": [(6, 30), (12, 45), (25, 90)],
                          "countersink": [(12, 20), (25, 35)],
                          "counterbore": [(12, 45), (25, 80)]}.items():
        for dia, secs in rows:
            t.add(MachiningRate(feature, float(dia), float(secs), 0.0, 20.0, 85.0, False, CALIBRATE_ME))
    t.add(MachiningRate("pocket", 1e9, 0.0, 15.0, 20.0, 85.0, False, CALIBRATE_ME))
    t.add(AssemblyRate("*", 10.0, 2.0, 0.5, 60.0, False, CALIBRATE_ME))
    return t
