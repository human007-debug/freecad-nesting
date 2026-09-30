"""
quoting/tender_costing.py
-------------------------
A tender costing sheet, the way it is done by hand in Excel -- but
computed from the BOM and a shop rate card (quoting/rate_card.py) instead
of typed:

  BOM items (consolidated bill of material: description, shape, grade,
  thickness / width / length, qty per set, qty per sheet or bar)
      -> net and gross weights by the sheet formula
      -> A. MATERIALS  grouped by shape, size and grade, at rate/kg x usage %
                       less scrap credit (gross - net) x scrap rate/kg
      -> B. PROCESS    per metre / stroke / piece / kg / sq ft; quantities
                       derived from the BOM where it can (laser metres from
                       part perimeters, cut counts, weights, fasteners from
                       the fastener price list), every one overridable
      -> C..H OTHERS   bought-out, S&I, packing, inspection, others, lab test
      -> I ALL TOTAL, J MARGIN (or CONTINGENCIES), K BASIC COST, L GST,
         M ALL INCLUSIVE RATE -- per unit (pair / set), and for the tender
         quantity
      -> LPR comparison: the margin that would make M equal the LPR, and
         a rounded QUOTED RATE.

The margin is the last knob, set per quote (`margin_pct`), the way a
costing is finished off; `margin_for_rate()` gives the margin that hits a
target all-inclusive rate.

`import_costing_workbook()` reads an existing Excel costing (the COSTING
SHEET + RAW MATERIAL layout) so past jobs can be re-costed against the
current rate card and compared line by line.
"""

import math
import re
from dataclasses import dataclass, field
from typing import Dict, List, Optional, Tuple

from quoting.rate_card import RateCard, RateCardError, grade_family, shape_class

KINDS = ("sheet", "plate", "tube", "bar", "fastener", "other")


def _num(v) -> Optional[float]:
    if v is None or v == "":
        return None
    if isinstance(v, (int, float)):
        return float(v)
    m = re.search(r"[-+]?\d*\.?\d+", str(v).replace(",", ""))
    return float(m.group()) if m else None


def _dims(v) -> List[float]:
    """'40*40' -> [40, 40]; 'Ø28*2mm' -> [28, 2]; '2x12x25mm' -> [2, 12, 25]; 28 -> [28]."""
    if v is None or v == "":
        return []
    if isinstance(v, (int, float)):
        return [float(v)]
    return [float(x) for x in re.findall(r"\d*\.?\d+", str(v).replace(",", ""))]


def _norm(s) -> str:
    return " ".join(str(s or "").strip().lower().replace("_", " ").split())


def _item_key(s) -> str:
    """Fastener/part names compared loosely: 'HEX. HD SCREW (M8*35)' ==
    'hex hd screw m8x35'."""
    t = str(s or "").lower().replace(".", " ")
    t = re.sub(r"(\d)\s*[x×*]\s*(\d)", r"\1x\2", t)
    t = re.sub(r"[()\[\],]", " ", t)
    return " ".join(t.split())


# ----------------------------------------------------------------- BOM

@dataclass
class BomItem:
    description: str
    shape: str                           # as written: SHEET, PLATE, SQ. TUBE, PIPE, ROD, FLAT, FASTNERS ...
    grade: str = ""
    thickness: Optional[float] = None    # mm (wall for a tube, diameter for a round rod)
    width: object = None                 # mm, or "40*40" for a section
    length: Optional[float] = None       # mm
    qty: float = 1.0                     # per set (all BOM levels multiplied)
    qty_per_stock: Optional[float] = None  # hand-counted qty per sheet / bar; None = formula
    stock_weight: Optional[float] = None   # kg of the stock piece (sheet, strip, bar) if not standard
    drawing: str = ""
    item: str = ""
    processes: Dict[str, float] = field(default_factory=dict)   # BOM process columns, per set

    @property
    def kind(self) -> str:
        s = _norm(self.shape)
        if "fast" in s or "bolt" in s or "screw" in s or "rivet" in s:
            return "fastener"
        cls = shape_class(self.shape)
        return cls if cls in KINDS else "other"


@dataclass
class ItemWeights:
    item: BomItem
    net_pc: float = 0.0
    gross_pc: float = 0.0
    qty_per_stock: Optional[float] = None
    note: str = ""

    @property
    def net(self) -> float:
        return self.net_pc * self.item.qty

    @property
    def gross(self) -> float:
        return self.gross_pc * self.item.qty


def item_weights(it: BomItem, card: RateCard) -> ItemWeights:
    f = card.formula
    w = ItemWeights(it)
    kind = it.kind
    t, L = it.thickness, it.length
    dims = _dims(it.width)
    if kind in ("fastener", "other") or not t or not L:
        return w
    if kind in ("sheet", "plate"):
        width = dims[0] if dims else None
        if not width:
            w.note = "no width"
            return w
        w.net_pc = f.piece_weight(t, width, L)
        if it.stock_weight and it.qty_per_stock:
            w.gross_pc = it.stock_weight / it.qty_per_stock
            w.qty_per_stock = it.qty_per_stock
        else:
            try:
                w.gross_pc = f.gross_weight(t, width, L, qty_per_sheet=int(it.qty_per_stock or 0) or None)
                w.qty_per_stock = it.qty_per_stock or f.qty_per_sheet(t, width, L)
            except ValueError as e:
                w.gross_pc, w.note = w.net_pc, str(e)
        return w
    if kind == "tube":
        name = _norm(it.shape)
        if len(dims) >= 2 and ("sq" in name or "rec" in name or "rhs" in name or "shs" in name):
            a, b = dims[-2], dims[-1]
            net = f.rect_tube_weight(t, a, b, L) if a != b else f.square_tube_weight(t, a, L)
            gross = f.rect_tube_weight(t, a, b, L + f.tube_cut_allowance) if a != b \
                else f.square_tube_weight(t, a, L + f.tube_cut_allowance)
        else:
            od = dims[0] if dims else None
            if not od:
                w.note = "no tube size"
                return w
            net = f.round_tube_weight(t, od, L)
            gross = f.round_tube_weight(t, od, L + f.tube_cut_allowance)
        w.net_pc = net
        if it.stock_weight and it.qty_per_stock:
            w.gross_pc, w.qty_per_stock = it.stock_weight / it.qty_per_stock, it.qty_per_stock
        else:
            w.gross_pc, w.qty_per_stock = gross, 1
        return w
    # bar: round rod (thickness = dia, no width), square rod / flat (t x w)
    area = (t * dims[0]) if dims else (math.pi / 4 * t * t)
    w.net_pc = area * L * f.kg_per_mm3
    bar_len = float(card.setting("bar_length") or 6000.0)
    per_bar = it.qty_per_stock or max(1, math.floor(bar_len / (L + 5)))
    stock = it.stock_weight or area * bar_len * f.kg_per_mm3
    w.gross_pc, w.qty_per_stock = stock / per_bar, per_bar
    return w


# ---------------------------------------------------------------- lines

@dataclass
class MaterialLine:
    shape: str
    size: str
    material: str
    weight: float                        # kg per unit (gross, as the sheets' "WEIGHT" column)
    factor: float                        # usage %
    rate: Optional[float]                # per kg
    remarks: str = ""
    net: float = 0.0
    scrap_rate: float = 0.0

    @property
    def charged_weight(self) -> float:
        return self.weight * self.factor

    @property
    def value(self) -> Optional[float]:
        return None if self.rate is None else self.charged_weight * self.rate


@dataclass
class ProcessLine:
    description: str
    uom: str
    qty: float
    rate: Optional[float]
    remarks: str = ""
    derived_qty: Optional[float] = None   # what the BOM suggested, when qty was typed over it

    @property
    def value(self) -> Optional[float]:
        return None if self.rate is None else self.qty * self.rate


@dataclass
class OtherLine:
    code: str
    item: str
    amount: float = 0.0
    remarks: str = ""


@dataclass
class TenderHeader:
    customer: str = ""
    division: str = ""
    tender_type: str = ""
    delivery_schedule: str = ""
    document_number: str = ""
    tender_number: str = ""
    pl_number: str = ""
    tender_due_date: str = ""
    product: str = ""
    drawing_number: str = ""
    quantity: float = 1.0
    unit: str = "Set"


FIELD_LABELS = (("customer", "CUSTOMER NAME"), ("division", "DIVISION"), ("tender_type", "TENDER TYPE"),
                ("delivery_schedule", "DELIVERY SCHEDULE"), ("document_number", "DOCUMENT NUMBER"),
                ("tender_number", "TENDER NUMBER"), ("pl_number", "PL NUMBER"),
                ("tender_due_date", "TENDER DUE DATE"), ("product", "PRODUCT"),
                ("drawing_number", "DRAWING NUMBER"))


# ---------------------------------------------------------------- costing

class TenderCosting:
    def __init__(self, card: RateCard, header: Optional[TenderHeader] = None,
                 items: Optional[List[BomItem]] = None):
        self.card = card
        self.header = header or TenderHeader()
        self.items: List[BomItem] = list(items or [])
        self.material_overrides: Dict[Tuple[str, str, str], dict] = {}   # key -> {"factor"/"rate"/"weight"}
        self.processes: Optional[List[ProcessLine]] = None   # None = derive from the BOM
        self.others: List[OtherLine] = [OtherLine(o.code, o.item, o.amount) for o in card.others]
        self.margin_pct: float = float(card.setting("margin_pct") or 0)
        self.margin_label: str = str(card.setting("margin_label") or "MARGIN")
        self.gst_pct: float = float(card.setting("gst_pct") or 0)
        self.scrap_kg_override: Optional[float] = None
        self.lpr: Optional[float] = None
        self.quoted_rate: Optional[float] = None
        self.notes: List[str] = []

    # -------------------------------------------------------------- weights

    def weights(self) -> List[ItemWeights]:
        return [item_weights(it, self.card) for it in self.items]

    @property
    def net_weight(self) -> float:
        return sum(w.net for w in self.weights())

    @property
    def gross_weight(self) -> float:
        return sum(w.gross for w in self.weights())

    # ------------------------------------------------------------ materials

    @staticmethod
    def _size_label(it: BomItem) -> str:
        k, t = it.kind, it.thickness
        dims = _dims(it.width)
        if k in ("sheet", "plate"):
            return f"{t:.2f}mm"
        if k == "tube":
            if len(dims) >= 2:
                a, b = dims[-2], dims[-1]
                return (f"{a:g}SQ*{t:g}mm" if a == b else f"{a:g}x{b:g}*{t:g}mm")
            return f"Ø{dims[0]:g}*{t:g}mm" if dims else f"{t:g}mm"
        if k == "bar":
            return f"{t:g}x{dims[0]:g}mm" if dims else f"Ø{t:g}mm"
        return ""

    def material_lines(self) -> List[MaterialLine]:
        groups: Dict[tuple, MaterialLine] = {}
        order = []
        for w in self.weights():
            it = w.item
            if it.kind not in ("sheet", "plate", "tube", "bar") or w.gross <= 0:
                continue
            fam = grade_family(it.grade) or ""
            key = (it.kind, self._size_label(it), fam)
            if key not in groups:
                try:
                    m = self.card.material(it.grade, it.kind, it.thickness)
                    rate, factor, scrap, remarks = m.rate_per_kg, m.usage_factor, m.scrap_per_kg, ""
                    label = m.material
                except RateCardError as e:
                    rate, factor, scrap, remarks, label = None, 1.0, 0.0, str(e), it.grade
                shape = {"sheet": "SHEET", "plate": "PLATE", "tube": it.shape.upper(), "bar": it.shape.upper()}
                groups[key] = MaterialLine(shape[it.kind], key[1], label, 0.0, factor, rate, remarks,
                                           scrap_rate=scrap)
                order.append(key)
            groups[key].weight += w.gross
            groups[key].net += w.net
        lines = []
        for key in order:
            line = groups[key]
            ov = self.material_overrides.get(key, {})
            for attr in ("factor", "rate", "weight"):
                if attr in ov:
                    setattr(line, attr, ov[attr])
            lines.append(line)
        return lines

    def scrap_kg(self) -> float:
        if self.scrap_kg_override is not None:
            return self.scrap_kg_override
        return max(self.gross_weight - self.net_weight, 0.0)

    def scrap_credit(self) -> float:
        """(gross - net) kg at each material's scrap rate, weighted by where the offcut comes from."""
        lines = self.material_lines()
        total_off = sum(max(l.weight - l.net, 0) for l in lines)
        if total_off <= 0:
            return 0.0
        kg = self.scrap_kg()
        return sum(kg * (max(l.weight - l.net, 0) / total_off) * l.scrap_rate for l in lines)

    @property
    def material_cost(self) -> float:
        return sum(l.value or 0.0 for l in self.material_lines()) - self.scrap_credit()

    # ------------------------------------------------------------ processes

    def derived_processes(self) -> List[ProcessLine]:
        """What the BOM says the process quantities are -- the starting point;
        a costing usually adds to laser metres for holes and slots."""
        card, out = self.card, []
        ws = self.weights()
        laser: Dict[float, float] = {}
        for w in ws:
            it = w.item
            if it.kind in ("sheet", "plate") and it.thickness and it.length and _dims(it.width):
                per = 2 * (_dims(it.width)[0] + it.length) * it.qty / 1000.0
                laser[it.thickness] = laser.get(it.thickness, 0.0) + per
        for t in sorted(laser):
            out.append(self._pline(f"LASER CUTTING {t:g}mm", "LASER CUTTING", "METRE", laser[t], t,
                                   "outer perimeters -- add holes & slots"))
        tube_cuts = sum(w.item.qty for w in ws if w.item.kind == "tube")
        if tube_cuts:
            out.append(self._pline("TUBE CUTTING", "TUBE CUTTING", "NOS", tube_cuts))
        rod_cuts = sum(w.item.qty for w in ws if w.item.kind == "bar")
        if rod_cuts:
            out.append(self._pline("ROD CUTTING", "ROD CUTTING", "NOS", rod_cuts))
        bom_totals: Dict[str, float] = {}
        for w in ws:
            for name, q in w.item.processes.items():
                if q:
                    bom_totals[name] = bom_totals.get(name, 0.0) + q
        for name, q in bom_totals.items():
            if _norm(name).startswith("laser") or _norm(name) in ("tube cutting", "rod cutting"):
                continue
            uom = next((p.uom for p in card.processes if _norm(p.process) == _norm(name)), "NOS")
            out.append(self._pline(name.upper(), name, uom, q, None, "from BOM"))
        fast_value, unmatched = self.fastener_value()
        if fast_value or unmatched:
            out.append(ProcessLine("FASTENERS", "SET", 1, fast_value,
                                   "priced from the fastener list" +
                                   (f"; not on the list: {', '.join(unmatched)}" if unmatched else "")))
        return out

    def _pline(self, label, name, uom, qty, thickness=None, remarks=""):
        try:
            rate = self.card.process_rate(name, thickness)
        except (RateCardError, ValueError):
            rate = None
            remarks = (remarks + "; " if remarks else "") + "no rate on the card"
        return ProcessLine(label, uom, round(qty, 3), rate, remarks)

    def fastener_value(self) -> Tuple[float, List[str]]:
        prices = {_item_key(f.item): f.price for f in self.card.fasteners}
        total, missing = 0.0, []
        for it in self.items:
            if it.kind != "fastener":
                continue
            key = _item_key(it.description)
            price = prices.get(key)
            if price is None:
                price = next((p for k, p in prices.items() if k and (k in key or key in k)), None)
            if price is None:
                missing.append(it.description)
            else:
                total += price * it.qty
        return total, missing

    def process_lines(self) -> List[ProcessLine]:
        return self.processes if self.processes is not None else self.derived_processes()

    @property
    def process_cost(self) -> float:
        return sum(p.value or 0.0 for p in self.process_lines())

    # --------------------------------------------------------------- totals

    @property
    def others_total(self) -> float:
        return sum(o.amount for o in self.others)

    @property
    def all_total(self) -> float:
        return self.material_cost + self.process_cost + self.others_total

    def totals(self, margin_pct: Optional[float] = None) -> dict:
        m = self.margin_pct if margin_pct is None else margin_pct
        total = self.all_total
        margin = total * m / 100.0
        basic = total + margin
        gst = basic * self.gst_pct / 100.0
        rate = basic + gst
        weight = self.gross_weight or 1.0
        return {
            "material": self.material_cost, "process": self.process_cost, "others": self.others_total,
            "all_total": total, "margin_pct": m, "margin": margin,
            "margin_share": margin / rate if rate else 0.0,
            "basic": basic, "gst": gst, "all_inclusive": rate,
            "material_per_kg": self.material_cost / weight, "process_per_kg": self.process_cost / weight,
            "tender_value": rate * (self.header.quantity or 1),
        }

    def margin_for_rate(self, all_inclusive: float) -> float:
        """The margin % that makes the ALL INCLUSIVE RATE equal `all_inclusive`."""
        total = self.all_total
        if total <= 0:
            raise ValueError("nothing to cost yet")
        basic = all_inclusive / (1 + self.gst_pct / 100.0)
        return (basic / total - 1) * 100.0

    def rounded_rate(self, rate: float) -> float:
        step = float(self.card.setting("quote_rounding") or 0)
        return round(rate / step) * step if step else rate


# ------------------------------------------------------------ Excel import

def _find_header(ws, wanted: str, rows=range(1, 8)):
    for r in rows:
        for c in range(1, min(ws.max_column, 120) + 1):
            if _norm(ws.cell(r, c).value) == wanted:
                return r, c
    return None


def read_bom_sheet(ws) -> List[BomItem]:
    """A CONSOLIDATED BILL OF MATERIAL worksheet (openpyxl, data_only):
    columns found by their headings, wherever they are."""
    heads: Dict[str, int] = {}
    for r in (3, 4):
        for c in range(1, min(ws.max_column, 120) + 1):
            name = _norm(ws.cell(r, c).value)
            if name and name not in heads:
                heads[name] = c
    col = lambda *names: next((heads[n] for n in names if n in heads), None)  # noqa: E731
    c_desc, c_type = col("description"), col("raw material spec", "type")
    c_grade, c_thk, c_w, c_l = col("grade"), col("thick"), col("width"), col("lnth", "length")
    c_qty, c_qps, c_sheet = col("total qty"), col("qty/sheet"), col("sheet wt.")
    c_drg = col("main drawing")
    process_cols = {}
    for name in ("bending (r.feets)", "bending", "drilling", "drill", "tapping", "notching", "piercing",
                 "tube bending", "machining", "chamfering"):
        if name in heads:
            process_cols[heads[name]] = name.split(" (")[0].upper()
    items = []
    for r in range(5, ws.max_row + 1):
        shape = ws.cell(r, c_type).value if c_type else None
        desc = ws.cell(r, c_desc).value if c_desc else None
        if not shape and not desc:
            continue
        qty = _num(ws.cell(r, c_qty).value) if c_qty else None
        if not qty:
            continue
        procs = {}
        for c, name in process_cols.items():
            v = _num(ws.cell(r, c).value)
            if v:
                procs[name] = procs.get(name, 0.0) + v * (1 if name else 1)
        items.append(BomItem(
            description=str(desc or shape).strip(), shape=str(shape or "").strip(),
            grade=str(ws.cell(r, c_grade).value or "").strip() if c_grade else "",
            thickness=_num(ws.cell(r, c_thk).value) if c_thk else None,
            width=ws.cell(r, c_w).value if c_w else None,
            length=_num(ws.cell(r, c_l).value) if c_l else None,
            qty=qty,
            qty_per_stock=_num(ws.cell(r, c_qps).value) if c_qps else None,
            stock_weight=_num(ws.cell(r, c_sheet).value) if c_sheet else None,
            drawing=str(ws.cell(r, c_drg).value or "") if c_drg else "",
            processes=procs,
        ))
    return items


def read_costing_sheet(ws) -> dict:
    """Header fields, process lines, others, margin/GST and LPR/quoted rate
    from a COSTING SHEET worksheet (openpyxl, data_only)."""
    out = {"header": TenderHeader(), "processes": [], "others": [], "margin_pct": None,
           "margin_label": None, "gst_pct": None, "lpr": None, "quoted": None, "materials": [],
           "totals": {}, "gross_weight": None}
    total_rows = {"TOTAL MATERIAL COST": "material", "TOTAL PROCESS COST": "process",
                  "ALL INCLUSIVE RATE": "all_inclusive"}
    labels = dict((lbl, attr) for attr, lbl in FIELD_LABELS)
    section = None
    for r in range(1, ws.max_row + 1):
        a, b = _norm(ws.cell(r, 1).value).upper(), _norm(ws.cell(r, 2).value).upper()
        if a in labels:
            v = ws.cell(r, 6).value
            setattr(out["header"], labels[a], v.date().isoformat() if hasattr(v, "date") else str(v or ""))
            continue
        if a == "QUANTITY IN NOS":
            v = str(ws.cell(r, 6).value or "")
            out["header"].quantity = _num(v) or 1.0
            out["header"].unit = (re.sub(r"[\d.\s]+", " ", v).strip() or "Set").split()[0].capitalize()
            continue
        if a in ("MATERIALS", "PROCESS", "OTHERS"):
            section = a
            continue
        if b in total_rows:
            out["totals"][total_rows[b]] = _num(ws.cell(r, 10).value)
            if b == "TOTAL MATERIAL COST":
                out["gross_weight"] = _num(ws.cell(r, 6).value)
            continue
        if b.startswith("ALL TOTAL"):
            out["totals"]["all_total"] = _num(ws.cell(r, 10).value)
            continue
        if b in ("MARGIN", "CONTINGENCIES"):
            out["margin_label"], out["margin_pct"] = b, (_num(ws.cell(r, 6).value) or 0) * 100
            continue
        if b == "GST":
            out["gst_pct"] = (_num(ws.cell(r, 6).value) or 0) * 100
            continue
        if b == "LPR":
            out["lpr"] = next((_num(ws.cell(r, c).value) for c in range(6, 13) if _num(ws.cell(r, c).value)), None)
            continue
        joined = " ".join(_norm(ws.cell(r, c).value) for c in range(1, 12))
        if "quoted rate" in joined and out["quoted"] is None:
            out["quoted"] = next((_num(ws.cell(r, c).value) for c in range(6, 13)
                                  if isinstance(ws.cell(r, c).value, (int, float))), None)
            continue
        if section == "PROCESS" and isinstance(ws.cell(r, 1).value, (int, float)) and b:
            qty = next((_num(ws.cell(r, c).value) for c in (8, 7) if _num(ws.cell(r, c).value) is not None), 0.0)
            out["processes"].append(ProcessLine(str(ws.cell(r, 2).value).strip(), str(ws.cell(r, 6).value or ""),
                                                qty, _num(ws.cell(r, 9).value)))
        if section == "MATERIALS" and isinstance(ws.cell(r, 1).value, (int, float)) and b:
            out["materials"].append({"shape": ws.cell(r, 2).value, "size": ws.cell(r, 3).value,
                                     "material": ws.cell(r, 4).value, "weight": _num(ws.cell(r, 6).value),
                                     "factor": _num(ws.cell(r, 7).value), "rate": _num(ws.cell(r, 9).value),
                                     "value": _num(ws.cell(r, 10).value)})
        if section == "OTHERS" and len(a) == 1 and a in "CDEFGH" and b:
            out["others"].append(OtherLine(a, str(ws.cell(r, 2).value).strip(),
                                           _num(ws.cell(r, 10).value) or 0.0))
    return out


def card_rate_for(card: RateCard, description: str, sheet_rate: Optional[float] = None):
    """(rate, remark) for a process line as a costing sheet words it --
    'LASER CUTTING(1.25mm)', 'PIERICNG', 'SQ. ROD & ROD MACHINING'. A
    thickness in the text picks the laser band; a process the card doesn't
    have (or a per-mm one with no thickness) keeps the sheet's own rate."""
    import difflib
    text = _norm(description)
    m = re.search(r"(?<![<>\d.])(\d+(?:\.\d+)?)\s*mm", text)
    thickness = float(m.group(1)) if m and float(m.group(1)) <= 50 else None
    base = re.sub(r"\(.*?\)|[\d.]+\s*mm", " ", text)
    base = " ".join(base.replace("(", " ").replace(")", " ").split())
    names = {_norm(p.process): p.process for p in card.processes}
    candidates = [text, base] + difflib.get_close_matches(base, list(names), n=1, cutoff=0.85)
    candidates += [n for n in sorted(names, key=len, reverse=True) if n in base]
    for cand in candidates:
        if cand not in names:
            continue
        try:
            return card.process_rate(names[cand], thickness), ""
        except (RateCardError, ValueError):
            continue
    if sheet_rate is not None:
        return sheet_rate, "not priced from the rate card -- the sheet's own rate is kept"
    return None, "no rate"


def import_costing_workbook(path: str, card: RateCard, costing_sheet: Optional[str] = None,
                            bom_sheet: Optional[str] = None) -> Tuple[TenderCosting, dict]:
    """A past Excel costing -> (TenderCosting on `card`, what the workbook
    itself said). The BOM is taken as-is (hand-counted qty/sheet and
    non-standard sheets included); processes keep the workbook's own
    quantities (they were counted by hand) but are re-priced from `card`."""
    import openpyxl
    wb = openpyxl.load_workbook(path, data_only=True)
    names = wb.sheetnames
    cs = wb[costing_sheet or next(n for n in names if _norm(n) == "costing sheet")]
    bs = wb[bom_sheet or next(n for n in names if _norm(n).startswith("raw material"))]
    sheet = read_costing_sheet(cs)
    tc = TenderCosting(card, sheet["header"], read_bom_sheet(bs))
    lines = []
    for p in sheet["processes"]:
        rate, how = card_rate_for(card, p.description, p.rate)
        lines.append(ProcessLine(p.description, p.uom, p.qty, rate, how))
    if lines:
        tc.processes = lines
    if sheet["others"]:
        tc.others = sheet["others"]
    if sheet["margin_pct"] is not None:
        tc.margin_pct = sheet["margin_pct"]
    if sheet["margin_label"]:
        tc.margin_label = sheet["margin_label"]
    if sheet["gst_pct"] is not None:
        tc.gst_pct = sheet["gst_pct"]
    tc.lpr, tc.quoted_rate = sheet["lpr"], sheet["quoted"]
    sheet["totals"]["others"] = sum(o.amount for o in sheet["others"])
    return tc, sheet
